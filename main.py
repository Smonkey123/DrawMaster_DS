import pdfplumber
import os
import PyPDF2
import fitz
import re
from openpyxl import Workbook, load_workbook, styles
import tkinter as tk
from tkinter import filedialog, messagebox, ttk


def extract_switchgear_number(title):
    """
    从标题中提取站号

    参数:
        title: 标题字符串

    返回:
        str: 站号（如 A01, A02...A99），如果不符合格式则返回None
    """
    if not title or len(title) < 5:
        return None

    first_5_chars = title[:5]
    match = re.match(r'^==([A-Z]\d{2})', first_5_chars)
    if match:
        return match.group(1)
    return None


def find_project_number(pdf_path, page_number=None):
    """
    在PDF图纸中查找CONTRACT NO.:字样，在其右下角找到项目号（9位数字）

    参数:
        pdf_path: PDF文件路径
        page_number: 指定页码（从0开始），如果为None则搜索所有页面

    返回:
        str: 找到的项目号（9位数字），如果未找到则返回None
    """
    doc = fitz.open(pdf_path)

    pages_to_search = [page_number] if page_number is not None else range(len(doc))

    for page_num in pages_to_search:
        if page_num >= len(doc):
            continue
        page = doc[page_num]

        search_patterns = ["CONTRACT NO.", "CONTRACT NO:", "CONTRACT NO"]

        for pattern in search_patterns:
            instances = page.search_for(pattern)

            if instances:
                bbox = instances[0]

                search_rect = fitz.Rect(
                    bbox.x1,
                    bbox.y0,
                    bbox.x1 + 200,
                    bbox.y1 + 50
                )

                text_in_area = page.get_text("text", clip=search_rect)
                numbers = re.findall(r'\b(\d{9})\b', text_in_area)

                if numbers:
                    project_number = numbers[0]
                    doc.close()
                    return project_number
                else:
                    expanded_rect = fitz.Rect(
                        bbox.x0 - 50,
                        bbox.y0 - 10,
                        bbox.x1 + 300,
                        bbox.y1 + 100
                    )
                    expanded_text = page.get_text("text", clip=expanded_rect)
                    numbers = re.findall(r'\b(\d{9})\b', expanded_text)
                    if numbers:
                        project_number = numbers[0]
                        doc.close()
                        return project_number

    doc.close()
    return None


def extract_valid_rectangles(pdf_path, page_number, search_terms, output_image_path=None):
    """
    从PDF页面的get_drawings()结果中提取包含所有搜索字符串的矩形区域

    参数:
        pdf_path: PDF文件路径
        page_number: 页码（从0开始）
        search_terms: 搜索字符串列表
        output_image_path: 输出标注图像的路径（可选）

    返回:
        list: 所有表格数据的列表，每个元素是一个表格的二维列表
    """
    string_bboxes = []
    all_tables_data = []

    try:
        doc = fitz.open(pdf_path)
        page = doc[page_number]
        reference_fonts = []

        # 先处理前两个字符串，获取参考字体
        if len(search_terms) >= 2:
            # 搜索第一个字符串（中文）
            first_term = search_terms[0]
            first_instances = page.search_for(first_term)
            # print(f"第一个字符串: '{first_term}', 实例: {first_instances}")

            # 搜索第二个字符串（英文）
            second_term = search_terms[1]
            second_instances = page.search_for(second_term)
            # print(f"第二个字符串: '{second_term}', 实例: {second_instances}")

            # 获取页面中的文本块
            blocks = page.get_text("dict", flags=fitz.TEXTFLAGS_TEXT)['blocks']

            # 提取参考字体
            for term, instances in [(first_term, first_instances), (second_term, second_instances)]:
                if instances:
                    for bbox in instances:
                        for block in blocks:
                            if 'lines' in block:
                                for line in block['lines']:
                                    for span in line['spans']:
                                        span_bbox = fitz.Rect(span['bbox'])
                                        if span_bbox.intersects(bbox):
                                            font_name = span['font']
                                            font_size = span['size']
                                            reference_fonts.append((font_name, font_size))
                                            # print(f"添加参考字体: {font_name}, {font_size}")
                                            break
                        if reference_fonts:
                            break
                if reference_fonts:
                    break

        # print(f"最终参考字体: {reference_fonts}")

        # 遍历每个搜索字符串组（中文和其后的英文视为同一个组）
        i = 0
        while i < len(search_terms):
            # 中文和英文视为同一个组
            chinese_term = search_terms[i]
            english_term = search_terms[i + 1] if i + 1 < len(search_terms) else None

            # 搜索中文
            chinese_instances = page.search_for(chinese_term)
            # 搜索英文
            english_instances = page.search_for(english_term) if english_term else []

            # 先尝试使用中文实例
            if chinese_instances:
                # 获取页面中的文本块
                blocks = page.get_text("dict", flags=fitz.TEXTFLAGS_TEXT)['blocks']

                # 存储符合条件的中文实例
                valid_instances = []

                # 遍历每个中文实例
                for bbox in chinese_instances:
                    # 遍历文本块，找到包含该边界框的文本
                    for block in blocks:
                        if 'lines' in block:
                            for line in block['lines']:
                                for span in line['spans']:
                                    span_bbox = fitz.Rect(span['bbox'])
                                    # 检查边界框是否相交
                                    if span_bbox.intersects(bbox):
                                        # 获取字体信息
                                        font_name = span['font']
                                        font_size = span['size']

                                        # 检查字体和大小是否与参考字体一致
                                        font_match = False
                                        for ref_font, ref_size in reference_fonts:
                                            # 比较字体名称（忽略大小写）和大小（允许小误差）
                                            if font_name.lower() == ref_font.lower() and abs(font_size - ref_size) < 1.0:
                                                font_match = True
                                                break

                                        if font_match:
                                            valid_instances.append((bbox, font_name, font_size))
                                            # print(f"中文字符串: '{chinese_term}', 字体: {font_name}, 大小: {font_size}, 边界框: {bbox} (匹配成功)")

                if valid_instances:
                    # 使用第一个符合条件的实例
                    selected_bbox = valid_instances[0][0]
                    string_bboxes.append((selected_bbox.x0, selected_bbox.y0, selected_bbox.x1, selected_bbox.y1))
                    # print(f"使用中文实例: 边界框: {selected_bbox}")
                else:
                    # 如果没有符合条件的中文实例，使用第一个中文实例
                    selected_bbox = chinese_instances[0]
                    string_bboxes.append((selected_bbox.x0, selected_bbox.y0, selected_bbox.x1, selected_bbox.y1))
                    # print(f"没有符合条件的中文实例，使用第一个中文实例: 边界框: {selected_bbox}")
            else:
                # 如果没有中文实例，使用英文实例
                if english_instances:
                    # 获取页面中的文本块
                    blocks = page.get_text("dict", flags=fitz.TEXTFLAGS_TEXT)['blocks']

                    # 存储符合条件的英文实例
                    valid_instances = []

                    # 遍历每个英文实例
                    for bbox in english_instances:
                        # 遍历文本块，找到包含该边界框的文本
                        for block in blocks:
                            if 'lines' in block:
                                for line in block['lines']:
                                    for span in line['spans']:
                                        span_bbox = fitz.Rect(span['bbox'])
                                        # 检查边界框是否相交
                                        if span_bbox.intersects(bbox):
                                            # 获取字体信息
                                            font_name = span['font']
                                            font_size = span['size']

                                            # 检查字体和大小是否与参考字体一致
                                            font_match = False
                                            for ref_font, ref_size in reference_fonts:
                                                # 比较字体名称（忽略大小写）和大小（允许小误差）
                                                if font_name.lower() == ref_font.lower() and abs(font_size - ref_size) < 1.0:
                                                    font_match = True
                                                    break

                                            if font_match:
                                                valid_instances.append((bbox, font_name, font_size))
                                                # print(f"英文字符串: '{english_term}', 字体: {font_name}, 大小: {font_size}, 边界框: {bbox} (匹配成功)")

                    if valid_instances:
                        # 使用第一个符合条件的实例
                        selected_bbox = valid_instances[0][0]
                        string_bboxes.append((selected_bbox.x0, selected_bbox.y0, selected_bbox.x1, selected_bbox.y1))
                        # print(f"使用英文实例: 边界框: {selected_bbox}")
                    else:
                        # 如果没有符合条件的英文实例，使用第一个英文实例
                        selected_bbox = english_instances[0]
                        string_bboxes.append((selected_bbox.x0, selected_bbox.y0, selected_bbox.x1, selected_bbox.y1))
                        # print(f"没有符合条件的英文实例，使用第一个英文实例: 边界框: {selected_bbox}")
                else:
                    print(f"未找到字符串 '{chinese_term}' 和 '{english_term}'")

            # 跳过下一个英文字符串
            i += 2
        # print(string_bboxes)
        # 从get_drawings()中获取图形元素
        drawings = page.get_drawings()

        # 首先找出所有符合基本条件的线条（不考虑宽度）
        candidate_lines_with_width = []
        if drawings:
            for drawing in drawings:
                # 检查是否满足基本条件（不考虑宽度）
                if (drawing.get('type') == 's' and
                        drawing.get('stroke_opacity') == 1.0 and
                        drawing.get('color') == (0.0, 0.0, 0.0) and
                        drawing.get('dashes') == '[] 0' and
                        drawing.get('lineCap') == (1, 1, 1)):

                    # 检查是否有items字段且包含线条
                    if 'items' in drawing:
                        for item in drawing['items']:
                            if isinstance(item, tuple) and item[0] == 'l':
                                # 这是一个线条
                                try:
                                    start_point = item[1]
                                    end_point = item[2]
                                    # 存储线条端点和宽度
                                    line_width = drawing.get('width', 0)
                                    candidate_lines_with_width.append((start_point, end_point, line_width))
                                except (AttributeError, IndexError):
                                    pass

        # 找出第一个字符串实例中心上方最近的线条的宽度
        target_width = 0.29180899262428284  # 默认宽度
        if string_bboxes and candidate_lines_with_width:
            # 获取第一个字符串实例的边界框
            string_left, string_top, string_right, string_bottom = string_bboxes[0]
            # 计算字符串中心坐标
            string_center_x = (string_left + string_right) / 2
            string_center_y = (string_top + string_bottom) / 2

            # 找出中心上方的水平线条
            upper_lines = []
            for start_point, end_point, line_width in candidate_lines_with_width:
                # 检查是否是水平线条
                is_horizontal = abs(start_point.y - end_point.y) < 0.1
                if is_horizontal:
                    line_y = start_point.y
                    # 检查是否在字符串中心上方
                    if line_y < string_center_y:
                        # 检查是否与字符串在水平方向上有重叠
                        line_min_x = min(start_point.x, end_point.x)
                        line_max_x = max(start_point.x, end_point.x)
                        has_overlap = not (line_max_x < string_left or line_min_x > string_right)
                        if has_overlap:
                            # 计算距离
                            distance = string_center_y - line_y
                            upper_lines.append((start_point, end_point, line_width, distance))

            # 找出最近的线条
            if upper_lines:
                upper_lines.sort(key=lambda x: x[3])
                nearest_line = upper_lines[0]
                target_width = nearest_line[2]
                # print(f"第一个字符串实例中心上方最近的线条宽度: {target_width}")

        # 使用目标宽度作为筛选条件，重新筛选线条
        all_lines = []
        if drawings:
            for drawing in drawings:
                # 检查是否满足特定条件
                if (drawing.get('type') == 's' and
                        drawing.get('stroke_opacity') == 1.0 and
                        drawing.get('color') == (0.0, 0.0, 0.0) and
                        abs(drawing.get('width', 0) - target_width) < 0.01 and  # 使用目标宽度作为筛选条件
                        drawing.get('dashes') == '[] 0' and
                        drawing.get('lineCap') == (1, 1, 1)):

                    # 检查是否有items字段且包含线条
                    if 'items' in drawing:
                        for item in drawing['items']:
                            if isinstance(item, tuple) and item[0] == 'l':
                                # 这是一个线条
                                try:
                                    start_point = item[1]
                                    end_point = item[2]
                                    # 存储线条端点
                                    all_lines.append((start_point, end_point))
                                except (AttributeError, IndexError):
                                    pass

        # 找出共同的端点
        common_points = []
        if all_lines:
            # 统计每个端点出现的次数
            point_counts = {}
            for start, end in all_lines:
                # 转换为可哈希的元组
                start_key = (round(start.x, 2), round(start.y, 2))
                end_key = (round(end.x, 2), round(end.y, 2))
                point_counts[start_key] = point_counts.get(start_key, 0) + 1
                point_counts[end_key] = point_counts.get(end_key, 0) + 1

            # 找出出现次数大于1的端点（共同端点）
            common_points = [fitz.Point(x, y) for (x, y), count in point_counts.items() if count > 1]

        # 生成候选矩形
        candidate_rectangles = []
        # 首先从共同端点生成候选矩形
        if len(common_points) >= 4:
            # 遍历所有可能的点组合
            for i in range(len(common_points)):
                for j in range(i + 1, len(common_points)):
                    p1 = common_points[i]
                    p2 = common_points[j]
                    # 确保p1在左上方，p2在右下方
                    left = min(p1.x, p2.x)
                    top = min(p1.y, p2.y)
                    right = max(p1.x, p2.x)
                    bottom = max(p1.y, p2.y)

                    # 检查该矩形是否包含所有字符串的中心
                    contains_all = True
                    for (string_left, string_top, string_right, string_bottom) in string_bboxes:
                        # 计算字符串的中心坐标
                        string_center_x = (string_left + string_right) / 2
                        string_center_y = (string_top + string_bottom) / 2
                        if not (left <= string_center_x <= right and top <= string_center_y <= bottom):
                            contains_all = False
                            break

                    if contains_all:
                        candidate_rectangles.append((left, top, right, bottom))

        # 如果从共同端点没有找到候选矩形，使用字符串边界框生成一个
        if not candidate_rectangles and string_bboxes:
            # 计算包含所有字符串的最小矩形
            min_x = min([bbox[0] for bbox in string_bboxes])
            min_y = min([bbox[1] for bbox in string_bboxes])
            max_x = max([bbox[2] for bbox in string_bboxes])
            max_y = max([bbox[3] for bbox in string_bboxes])
            candidate_rectangles.append((min_x, min_y, max_x, max_y))

        # 找出面积最小的矩形
        if candidate_rectangles:
            min_area = float('inf')
            best_rectangle = None
            for rect in candidate_rectangles:
                width = rect[2] - rect[0]
                height = rect[3] - rect[1]
                area = width * height
                if area < min_area:
                    min_area = area
                    best_rectangle = rect

            merged_bounding_box = best_rectangle
            # print(f"找到包含所有字符串的最小矩形: {merged_bounding_box}, 面积: {min_area}")
        else:
            # 如果没有找到候选矩形，使用整个页面
            merged_bounding_box = (0, 0, page.rect.width, page.rect.height)
            print("未找到包含所有字符串的矩形，使用整个页面作为区域")

        # 如果需要输出标注图像
        if output_image_path:
            # print(f"正在生成矩形标注图像: {output_image_path}")

            # 确保输出文件夹存在
            output_dir = os.path.dirname(output_image_path)
            if not os.path.exists(output_dir):
                os.makedirs(output_dir)
                # print(f"创建输出文件夹: {output_dir}")

            # 在页面上绘制标注
            if string_bboxes:
                # 绘制每个字符串的边界框（中文和英文合并后的边界框）
                for i, (left, top, right, bottom) in enumerate(string_bboxes):
                    shape = page.new_shape()
                    shape.draw_rect(fitz.Rect(left, top, right, bottom))
                    shape.finish(color=(0, 1, 0), fill=None, width=2)
                    shape.commit()
                    # # 绘制组编号
                    # shape = page.new_shape()
                    # shape.insert_text(fitz.Point(left + 5, top - 5), f"Group {i+1}", fontsize=10, color=(0, 1, 0))
                    # shape.commit()

                # 绘制共同端点（使用蓝色圆点标注）
                for point in common_points:
                    shape = page.new_shape()
                    # 绘制一个小圆点
                    shape.draw_circle(point, 4)  # 半径为4
                    shape.finish(color=(0, 0, 1), fill=(0, 0, 1), width=1)
                    shape.commit()

                # 绘制找到的最小矩形
                # print(f"绘制最小矩形: {merged_bounding_box}")
                shape = page.new_shape()
                shape.draw_rect(fitz.Rect(merged_bounding_box[0], merged_bounding_box[1], merged_bounding_box[2], merged_bounding_box[3]))
                # 使用更明显的颜色和线宽
                shape.finish(color=(1, 0, 0), fill=None, width=1)
                shape.commit()
                # print("最小矩形绘制完成")

                # 找出每个字符串上方的最近的线条（从已筛选的线条all_lines中）
                string_upper_lines = []
                if string_bboxes and all_lines:
                    # print("从all_lines中找出每个字符串上方的最近的线条...")

                    # 遍历每个字符串
                    for i, (string_left, string_top, string_right, string_bottom) in enumerate(string_bboxes):
                        # 存储候选线条
                        candidate_lines = []

                        # 遍历已筛选的线条
                        for start_point, end_point in all_lines:
                            # 检查是否为水平线条（y坐标差异很小）
                            is_horizontal = abs(start_point.y - end_point.y) < 0.1
                            if not is_horizontal:
                                continue

                            # 计算线条的坐标范围
                            line_min_x = min(start_point.x, end_point.x)
                            line_max_x = max(start_point.x, end_point.x)
                            line_y = start_point.y  # 水平线条的y坐标

                            # 计算字符串的中心y坐标
                            string_center_y = (string_top + string_bottom) / 2

                            # 检查线条是否在字符串的上方
                            is_above = line_y < string_center_y

                            # 检查线条是否与字符串在水平方向上有重叠
                            has_overlap = not (line_max_x < string_left or line_min_x > string_right)

                            if is_above and has_overlap:
                                # 计算线条与字符串中心的距离（垂直方向）
                                distance = string_center_y - line_y
                                candidate_lines.append((line_min_x, line_y, line_max_x, line_y, distance))

                        # 找到距离最近的线条
                        if candidate_lines:
                            # 按距离排序
                            candidate_lines.sort(key=lambda x: x[4])
                            # 选择最近的一个
                            nearest_line = candidate_lines[0][:4]
                            string_upper_lines.append(nearest_line)
                            # print(f"字符串 {i + 1} 上方的最近线条: {nearest_line}")

                            # 在标注图像中绘制这个线条
                            shape = page.new_shape()
                            shape.draw_line(fitz.Point(nearest_line[0], nearest_line[1]), fitz.Point(nearest_line[2], nearest_line[3]))
                            shape.finish(color=(0, 1, 0), width=2)
                            shape.commit()
                        else:
                            string_upper_lines.append(None)
                            # print(f"字符串 {i + 1} 上方未找到线条")

                # 找出merged_bounding_box内被线条分割的矩形区域
                if all_lines and merged_bounding_box:
                    print("找出被线条分割的矩形区域...")

                    # 1. 找出所有线条之间的交点
                    intersections = []

                    # 定义线段相交检测函数
                    def line_intersection(line1, line2):
                        """
                        检测两条线段是否相交，并返回交点
                        line1: ((x1, y1), (x2, y2))
                        line2: ((x3, y3), (x4, y4))
                        """
                        (x1, y1), (x2, y2) = line1
                        (x3, y3), (x4, y4) = line2

                        # 计算分母
                        denom = (y4 - y3) * (x2 - x1) - (x4 - x3) * (y2 - y1)
                        if denom == 0:
                            # 线条平行，检查是否重合或端点相交
                            # 检查line1的端点是否在线段line2上
                            if point_on_segment((x1, y1), (x3, y3), (x4, y4)):
                                return (x1, y1)
                            if point_on_segment((x2, y2), (x3, y3), (x4, y4)):
                                return (x2, y2)
                            # 检查line2的端点是否在线段line1上
                            if point_on_segment((x3, y3), (x1, y1), (x2, y2)):
                                return (x3, y3)
                            if point_on_segment((x4, y4), (x1, y1), (x2, y2)):
                                return (x4, y4)
                            return None

                        # 计算参数
                        ua = ((x4 - x3) * (y1 - y3) - (y4 - y3) * (x1 - x3)) / denom
                        ub = ((x2 - x1) * (y1 - y3) - (y2 - y1) * (x1 - x3)) / denom

                        # 检查交点是否在线段范围内
                        if 0 <= ua <= 1 and 0 <= ub <= 1:
                            x = x1 + ua * (x2 - x1)
                            y = y1 + ua * (y2 - y1)
                            return (x, y)
                        return None

                    def point_on_segment(point, line_start, line_end):
                        """
                        检查点是否在线段上
                        """
                        x, y = point
                        x1, y1 = line_start
                        x2, y2 = line_end

                        # 检查点是否在线段的边界框内
                        if not (min(x1, x2) - 0.001 <= x <= max(x1, x2) + 0.001 and
                                min(y1, y2) - 0.001 <= y <= max(y1, y2) + 0.001):
                            return False

                        # 检查点是否在直线上（斜率计算）
                        if abs(x2 - x1) < 0.001:
                            # 垂直线
                            return abs(x - x1) < 0.001
                        else:
                            # 计算斜率
                            slope = (y2 - y1) / (x2 - x1)
                            return abs(y - y1 - slope * (x - x1)) < 0.001

                    # 转换所有线条为端点元组
                    line_segments = []
                    for start, end in all_lines:
                        line_segments.append(((start.x, start.y), (end.x, end.y)))

                    # 检测所有线条对之间的交点
                    for i in range(len(line_segments)):
                        for j in range(i + 1, len(line_segments)):
                            line1 = line_segments[i]
                            line2 = line_segments[j]
                            intersection = line_intersection(line1, line2)
                            if intersection:
                                # 四舍五入到两位小数，避免浮点数精度问题
                                rounded_intersection = (round(intersection[0], 2), round(intersection[1], 2))
                                # 避免重复添加相同的交点
                                if rounded_intersection not in intersections:
                                    intersections.append(rounded_intersection)

                    # print(f"找到 {len(intersections)} 个线条交点")

                    # 在标注图像中绘制这些交点
                    for intersection in intersections:
                        shape = page.new_shape()
                        # 绘制一个小圆点表示交点，半径为2
                        shape.draw_circle(fitz.Point(intersection[0], intersection[1]), 2)
                        shape.finish(color=(1, 0, 0), fill=(1, 0, 0), width=1)
                        shape.commit()

                    # 检查交点数量是否为15个或16个
                    if len(intersections) not in [15, 16]:
                        print(f"警告：线条交点数量为 {len(intersections)}，不是预期的15个或16个！")

                    # 2. 生成所有可能的矩形组合（四个交点构成）
                    candidate_regions = []
                    # 遍历所有可能的四个交点组合
                    for i in range(len(intersections)):
                        for j in range(len(intersections)):
                            for k in range(len(intersections)):
                                for l in range(len(intersections)):
                                    # 确保四个点是不同的
                                    if i != j and i != k and i != l and j != k and j != l and k != l:
                                        p1 = intersections[i]
                                        p2 = intersections[j]
                                        p3 = intersections[k]
                                        p4 = intersections[l]

                                        # 计算边界
                                        all_x = [p1[0], p2[0], p3[0], p4[0]]
                                        all_y = [p1[1], p2[1], p3[1], p4[1]]
                                        left = min(all_x)
                                        top = min(all_y)
                                        right = max(all_x)
                                        bottom = max(all_y)

                                        # 检查是否形成有效的矩形
                                        if right > left and bottom > top:
                                            # 检查四个角是否都是交点
                                            corners = [(left, top), (right, top), (left, bottom), (right, bottom)]
                                            all_corners_are_intersections = True
                                            for corner in corners:
                                                # 检查角点是否在交点列表中（考虑浮点数精度）
                                                found = False
                                                for intersection in intersections:
                                                    if abs(corner[0] - intersection[0]) < 0.01 and abs(corner[1] - intersection[1]) < 0.01:
                                                        found = True
                                                        break
                                                if not found:
                                                    all_corners_are_intersections = False
                                                    break

                                            if all_corners_are_intersections:
                                                # 3. 检查是否在merged_bounding_box内
                                                c_left, c_top, c_right, c_bottom = merged_bounding_box
                                                if left >= c_left and top >= c_top and right <= c_right and bottom <= c_bottom:
                                                    # 避免重复添加相同的矩形
                                                    rectangle = (left, top, right, bottom)
                                                    if rectangle not in candidate_regions:
                                                        candidate_regions.append(rectangle)

                    # print(f"生成 {len(candidate_regions)} 个候选矩形区域")

                    # # 绘制所有候选矩形区域（用黑线）
                    # for region in candidate_regions:
                    #     r_left, r_top, r_right, r_bottom = region
                    #     shape = page.new_shape()
                    #     shape.draw_rect(fitz.Rect(r_left, r_top, r_right, r_bottom))
                    #     shape.finish(color=(0, 0, 0), fill=None, width=1)
                    #     shape.commit()
                    # print(f"绘制了 {len(candidate_regions)} 个候选矩形区域")

                    # 在merged_bounding_box中检测点状横线
                    all_dotted_lines = []
                    try:
                        drawings = page.get_drawings(extended=True)
                        for drawing in drawings:
                            if drawing.get('type') == 's':
                                dashes = drawing.get('dashes', None)
                                if dashes and isinstance(dashes, str) and len(dashes) > 2 and dashes.startswith('['):
                                    try:
                                        dash_pattern = dashes.split('[')[1].split(']')[0].strip()
                                        parts = dash_pattern.split()
                                        if len(parts) >= 2:
                                            dot_length = float(parts[0])
                                            if dot_length < 1.0:
                                                if 'rect' in drawing:
                                                    rect = drawing['rect']
                                                    try:
                                                        left = rect.x0 if hasattr(rect, 'x0') else rect[0]
                                                        top = rect.y0 if hasattr(rect, 'y0') else rect[1]
                                                        right = rect.x1 if hasattr(rect, 'x1') else rect[2]
                                                        bottom = rect.y1 if hasattr(rect, 'y1') else rect[3]

                                                        # 检查是否在merged_bounding_box内
                                                        c_left, c_top, c_right, c_bottom = merged_bounding_box
                                                        if (top >= c_top - 5 and top <= c_bottom + 5):
                                                            # 检查线段是否在水平方向上与merged_bounding_box有重叠
                                                            line_width = right - left
                                                            if line_width > 0:
                                                                # 计算线段的中心位置
                                                                line_center = (left + right) / 2
                                                                # 计算merged_bounding_box的中心位置
                                                                box_center = (c_left + c_right) / 2
                                                                # 计算水平偏移量
                                                                offset = box_center - line_center
                                                                # 平移线段，使其中心与merged_bounding_box的中心对齐
                                                                translated_left = left + offset
                                                                translated_right = right + offset
                                                                # 确保平移后的线段完全在merged_bounding_box内
                                                                if translated_left < c_left:
                                                                    # 如果左端点仍在bounding box外，向右平移
                                                                    offset = c_left - translated_left
                                                                    translated_left += offset
                                                                    translated_right += offset
                                                                elif translated_right > c_right:
                                                                    # 如果右端点仍在bounding box外，向左平移
                                                                    offset = c_right - translated_right
                                                                    translated_left += offset
                                                                    translated_right += offset
                                                                all_dotted_lines.append((translated_left, top, translated_right, bottom))
                                                            else:
                                                                # 线段宽度为0，直接添加
                                                                all_dotted_lines.append((left, top, right, bottom))
                                                    except (AttributeError, IndexError):
                                                        pass
                                    except (ValueError, IndexError):
                                        pass
                    except Exception as e:
                        print(f"检测点状横线时出错: {e}")

                    # print(f"在merged_bounding_box内检测到 {len(all_dotted_lines)} 个点状横线")

                    # 合并重合的点状横线（距离小于0.5的视为同一条）
                    merged_dotted_lines = []
                    if all_dotted_lines:
                        # 按y坐标排序
                        sorted_dotted_lines = sorted(all_dotted_lines, key=lambda x: x[1])

                        # 合并距离很近的横线
                        current_line = sorted_dotted_lines[0]
                        for line in sorted_dotted_lines[1:]:
                            _, current_y, _, _ = current_line
                            _, line_y, _, _ = line
                            if abs(line_y - current_y) < 0.5:
                                # 合并这两条线，取平均y坐标
                                merged_y = (current_y + line_y) / 2
                                merged_left = min(current_line[0], line[0])
                                merged_right = max(current_line[2], line[2])
                                current_line = (merged_left, merged_y, merged_right, merged_y)
                            else:
                                merged_dotted_lines.append(current_line)
                                current_line = line
                        # 添加最后一条线
                        merged_dotted_lines.append(current_line)

                    # print(f"合并后检测到 {len(merged_dotted_lines)} 个点状横线")

                    # 输出合并后的点状横线坐标，按y坐标排序
                    # print("合并后的点状横线坐标（按y坐标排序）:")
                    sorted_merged_lines = sorted(merged_dotted_lines, key=lambda x: x[1])
                    # for i, (left, top, right, bottom) in enumerate(sorted_merged_lines):
                    #     print(f"{i+1}. y={top:.2f}, left={left:.2f}, right={right:.2f}")

                    # 使用合并后的点状横线
                    all_dotted_lines = merged_dotted_lines

                    # 在标注图像中绘制所有点状横线（使用黄色）
                    for left, top, right, bottom in all_dotted_lines:
                        # 将端点限制在merged_bounding_box内
                        c_left, c_top, c_right, c_bottom = merged_bounding_box
                        constrained_left = max(left, c_left)
                        constrained_right = min(right, c_right)
                        shape = page.new_shape()
                        shape.draw_line(fitz.Point(constrained_left, top), fitz.Point(constrained_right, top))
                        shape.finish(color=(1, 1, 0), width=4)
                        shape.commit()

                    # 4. 找出包含对应字符串的矩形
                    string_regions = []
                    for i, (string_left, string_top, string_right, string_bottom) in enumerate(string_bboxes):
                        # 存储包含该字符串的矩形
                        containing_regions = []

                        for region in candidate_regions:
                            r_left, r_top, r_right, r_bottom = region
                            # 计算字符串的中心坐标
                            string_center_x = (string_left + string_right) / 2
                            string_center_y = (string_top + string_bottom) / 2
                            # 检查字符串中心是否在矩形内
                            if r_left <= string_center_x <= r_right and r_top <= string_center_y <= r_bottom:
                                containing_regions.append(region)

                        if containing_regions:
                            # 选择面积最小的矩形
                            min_area = float('inf')
                            best_region = None
                            for region in containing_regions:
                                width = region[2] - region[0]
                                height = region[3] - region[1]
                                area = width * height
                                if area < min_area:
                                    min_area = area
                                    best_region = region

                            string_regions.append(best_region)
                            # print(f"字符串 {i + 1} 所在的矩形区域: {best_region}")

                            # 找出该区域内的点状横线
                            region_dotted_lines = []
                            for dotted_line in all_dotted_lines:
                                left, top, right, bottom = dotted_line
                                r_left, r_top, r_right, r_bottom = best_region
                                if (top >= r_top - 5 and top <= r_bottom + 5):
                                    # 即使端点超出，只要在垂直方向上有重叠，就算作该区域内的点状横线
                                    region_dotted_lines.append(dotted_line)
                            # print(f"字符串 {i + 1} 所在区域内有 {len(region_dotted_lines)} 个点状横线")

                            # 定义每个字符串区域的竖向分割线规则
                            vertical_dividers = []
                            r_left, r_top, r_right, r_bottom = best_region
                            region_width = r_right - r_left

                            if i == 0:  # 第一个字符串区域
                                # 竖向分割线位置在表格横向的39.5/127和66/127
                                div1 = r_left + (39.5 / 127) * region_width
                                div2 = r_left + (66 / 127) * region_width
                                vertical_dividers = [div1, div2]
                            elif i in [1, 2]:  # 第二个和第三个字符串区域
                                # 竖向分割线位置在表格横向的66/127
                                div = r_left + (66 / 127) * region_width
                                vertical_dividers = [div]
                            elif i == 3:  # 第四个字符串区域
                                # 竖向分割线位置在表格横向的40.8/127、55.7/127、73.8/127、92/127
                                div1 = r_left + (40.8 / 127) * region_width
                                div2 = r_left + (55.7 / 127) * region_width
                                div3 = r_left + (73.8 / 127) * region_width
                                div4 = r_left + (92 / 127) * region_width
                                vertical_dividers = [div1, div2, div3, div4]
                            elif i == 4:  # 第五个字符串区域
                                # 竖向分割线位置在表格横向的55.7/127
                                div = r_left + (55.7 / 127) * region_width
                                vertical_dividers = [div]
                            elif i == 5:  # 第六个字符串区域
                                # 竖向分割线位置在表格横向的35/127
                                div = r_left + (35 / 127) * region_width
                                vertical_dividers = [div]
                            elif i == 6:  # 第七个字符串区域
                                # 竖向分割线位置有0根
                                vertical_dividers = []

                            # 绘制竖向分割线
                            for divider in vertical_dividers:
                                shape = page.new_shape()
                                shape.draw_line(fitz.Point(divider, r_top), fitz.Point(divider, r_bottom))
                                shape.finish(color=(0, 0, 1), width=1.5)
                                shape.commit()

                            # 准备表格分割
                            # 1. 收集所有水平分割线（点状横线）
                            horizontal_lines = []
                            for line in region_dotted_lines:
                                _, y, _, _ = line
                                horizontal_lines.append(y)
                            # 添加区域的上下边界
                            horizontal_lines.append(r_top)
                            horizontal_lines.append(r_bottom)
                            # 去重并排序
                            horizontal_lines = sorted(list(set(horizontal_lines)))

                            # 2. 收集所有垂直分割线
                            vertical_lines = [r_left] + vertical_dividers + [r_right]
                            # 排序
                            vertical_lines = sorted(vertical_lines)

                            # 3. 提取表格数据
                            table_data = []
                            min_row_height = 2  # 最小行高阈值

                            # 先过滤掉高度很小的行，收集有效的行索引
                            valid_rows = []
                            for row_idx in range(len(horizontal_lines) - 1):
                                top = horizontal_lines[row_idx]
                                bottom = horizontal_lines[row_idx + 1]
                                row_height = bottom - top

                                if row_height >= min_row_height:
                                    valid_rows.append((row_idx, top, bottom, row_height))

                            total_valid_rows = len(valid_rows)

                            for valid_row_idx, (original_row_idx, top, bottom, row_height) in enumerate(valid_rows):
                                row_data = []

                                # 处理第1个字符串矩形（i=0）的特殊合并规则
                                if i == 0:
                                    # 检查是否是最后两行（基于有效行数）
                                    is_last_two_rows = valid_row_idx >= total_valid_rows - 2

                                    if not is_last_two_rows:
                                        # 除了最后两行，其余的行前两列单元格合并提取
                                        if len(vertical_lines) >= 3:
                                            # 合并前两列
                                            left1 = vertical_lines[0]
                                            right1 = vertical_lines[2]  # 合并到第二列的右边界

                                            # 提取合并后的单元格文本
                                            cell_bbox = fitz.Rect(left1, top, right1, bottom)
                                            words = page.get_text("words", clip=cell_bbox)
                                            merged_text = ""
                                            if words:
                                                # 根据y坐标进行聚类，将错落范围小的单词视为同一行
                                                word_info = []
                                                for word in words:
                                                    x0, y0, x1, y1, text, block_no, line_no, word_no = word
                                                    center_y = (y0 + y1) / 2
                                                    word_info.append({
                                                        'text': text,
                                                        'x0': x0,
                                                        'x1': x1,
                                                        'center_y': center_y,
                                                        'height': y1 - y0
                                                    })

                                                # 按y坐标排序
                                                word_info.sort(key=lambda x: x['center_y'])

                                                # 聚类：将y坐标相近的单词归为同一行
                                                lines = []
                                                if word_info:
                                                    avg_height = sum([w['height'] for w in word_info]) / len(word_info)
                                                    line_threshold = avg_height * 0.5

                                                    current_line = [word_info[0]]
                                                    current_y = word_info[0]['center_y']

                                                    for word in word_info[1:]:
                                                        if abs(word['center_y'] - current_y) <= line_threshold:
                                                            current_line.append(word)
                                                        else:
                                                            lines.append(current_line)
                                                            current_line = [word]
                                                            current_y = word['center_y']
                                                    lines.append(current_line)

                                                # 对每一行的单词按x坐标排序
                                                for line in lines:
                                                    line.sort(key=lambda x: x['x0'])

                                                # 组合文本
                                                for line in lines:
                                                    line_text = " ".join([w['text'] for w in line])
                                                    if merged_text:
                                                        merged_text += "\n" + line_text
                                                    else:
                                                        merged_text = line_text
                                                merged_text = merged_text.strip()

                                            # 添加合并后的文本作为第一列
                                            row_data.append(merged_text)

                                            # 处理剩余列
                                            for col_idx in range(2, len(vertical_lines) - 1):
                                                left = vertical_lines[col_idx]
                                                right = vertical_lines[col_idx + 1]

                                                # 提取该单元格内的文本
                                                cell_bbox = fitz.Rect(left, top, right, bottom)
                                                words = page.get_text("words", clip=cell_bbox)
                                                cell_text = ""
                                                if words:
                                                    word_info = []
                                                    for word in words:
                                                        x0, y0, x1, y1, text, block_no, line_no, word_no = word
                                                        center_y = (y0 + y1) / 2
                                                        word_info.append({
                                                            'text': text,
                                                            'x0': x0,
                                                            'x1': x1,
                                                            'center_y': center_y,
                                                            'height': y1 - y0
                                                        })

                                                    word_info.sort(key=lambda x: x['center_y'])

                                                    lines = []
                                                    if word_info:
                                                        avg_height = sum([w['height'] for w in word_info]) / len(word_info)
                                                        line_threshold = avg_height * 0.5

                                                        current_line = [word_info[0]]
                                                        current_y = word_info[0]['center_y']

                                                        for word in word_info[1:]:
                                                            if abs(word['center_y'] - current_y) <= line_threshold:
                                                                current_line.append(word)
                                                            else:
                                                                lines.append(current_line)
                                                                current_line = [word]
                                                                current_y = word['center_y']
                                                        lines.append(current_line)

                                                    for line in lines:
                                                        line.sort(key=lambda x: x['x0'])

                                                    for line in lines:
                                                        line_text = " ".join([w['text'] for w in line])
                                                        if cell_text:
                                                            cell_text += "\n" + line_text
                                                        else:
                                                            cell_text = line_text
                                                    cell_text = cell_text.strip()

                                                row_data.append(cell_text)
                                    else:
                                        # 最后两行，按照原来的提取规则
                                        for col_idx in range(len(vertical_lines) - 1):
                                            left = vertical_lines[col_idx]
                                            right = vertical_lines[col_idx + 1]

                                            # 提取该单元格内的文本
                                            cell_bbox = fitz.Rect(left, top, right, bottom)
                                            words = page.get_text("words", clip=cell_bbox)
                                            cell_text = ""
                                            if words:
                                                word_info = []
                                                for word in words:
                                                    x0, y0, x1, y1, text, block_no, line_no, word_no = word
                                                    center_y = (y0 + y1) / 2
                                                    word_info.append({
                                                        'text': text,
                                                        'x0': x0,
                                                        'x1': x1,
                                                        'center_y': center_y,
                                                        'height': y1 - y0
                                                    })

                                                word_info.sort(key=lambda x: x['center_y'])

                                                lines = []
                                                if word_info:
                                                    avg_height = sum([w['height'] for w in word_info]) / len(word_info)
                                                    line_threshold = avg_height * 0.5

                                                    current_line = [word_info[0]]
                                                    current_y = word_info[0]['center_y']

                                                    for word in word_info[1:]:
                                                        if abs(word['center_y'] - current_y) <= line_threshold:
                                                            current_line.append(word)
                                                        else:
                                                            lines.append(current_line)
                                                            current_line = [word]
                                                            current_y = word['center_y']
                                                    lines.append(current_line)

                                                for line in lines:
                                                    line.sort(key=lambda x: x['x0'])

                                                for line in lines:
                                                    line_text = " ".join([w['text'] for w in line])
                                                    if cell_text:
                                                        cell_text += "\n" + line_text
                                                    else:
                                                        cell_text = line_text
                                                cell_text = cell_text.strip()

                                            row_data.append(cell_text)
                                # 处理第4个字符串矩形（i=3）的特殊合并规则
                                elif i == 3:
                                    # 第13行单元格合并提取（基于有效行数，第13行是valid_row_idx=12）
                                    if valid_row_idx == 12:
                                        # 合并所有列
                                        if len(vertical_lines) >= 2:
                                            left_all = vertical_lines[0]
                                            right_all = vertical_lines[-1]

                                            # 提取合并后的单元格文本
                                            cell_bbox = fitz.Rect(left_all, top, right_all, bottom)
                                            words = page.get_text("words", clip=cell_bbox)
                                            merged_text = ""
                                            if words:
                                                word_info = []
                                                for word in words:
                                                    x0, y0, x1, y1, text, block_no, line_no, word_no = word
                                                    center_y = (y0 + y1) / 2
                                                    word_info.append({
                                                        'text': text,
                                                        'x0': x0,
                                                        'x1': x1,
                                                        'center_y': center_y,
                                                        'height': y1 - y0
                                                    })

                                                word_info.sort(key=lambda x: x['center_y'])

                                                lines = []
                                                if word_info:
                                                    avg_height = sum([w['height'] for w in word_info]) / len(word_info)
                                                    line_threshold = avg_height * 0.5

                                                    current_line = [word_info[0]]
                                                    current_y = word_info[0]['center_y']

                                                    for word in word_info[1:]:
                                                        if abs(word['center_y'] - current_y) <= line_threshold:
                                                            current_line.append(word)
                                                        else:
                                                            lines.append(current_line)
                                                            current_line = [word]
                                                            current_y = word['center_y']
                                                    lines.append(current_line)

                                                for line in lines:
                                                    line.sort(key=lambda x: x['x0'])

                                                for line in lines:
                                                    line_text = " ".join([w['text'] for w in line])
                                                    if merged_text:
                                                        merged_text += "\n" + line_text
                                                    else:
                                                        merged_text = line_text
                                                merged_text = merged_text.strip()

                                            # 添加合并后的文本作为第一列
                                            row_data.append(merged_text)
                                            # 其余列留空
                                            for _ in range(1, len(vertical_lines) - 1):
                                                row_data.append("")
                                    # 第1行前两列合并提取（valid_row_idx=0）
                                    elif valid_row_idx == 0:
                                        if len(vertical_lines) >= 3:
                                            # 合并前两列
                                            left1 = vertical_lines[0]
                                            right1 = vertical_lines[2]  # 合并到第二列的右边界

                                            # 提取合并后的单元格文本
                                            cell_bbox = fitz.Rect(left1, top, right1, bottom)
                                            words = page.get_text("words", clip=cell_bbox)
                                            merged_text = ""
                                            if words:
                                                word_info = []
                                                for word in words:
                                                    x0, y0, x1, y1, text, block_no, line_no, word_no = word
                                                    center_y = (y0 + y1) / 2
                                                    word_info.append({
                                                        'text': text,
                                                        'x0': x0,
                                                        'x1': x1,
                                                        'center_y': center_y,
                                                        'height': y1 - y0
                                                    })

                                                word_info.sort(key=lambda x: x['center_y'])

                                                lines = []
                                                if word_info:
                                                    avg_height = sum([w['height'] for w in word_info]) / len(word_info)
                                                    line_threshold = avg_height * 0.5

                                                    current_line = [word_info[0]]
                                                    current_y = word_info[0]['center_y']

                                                    for word in word_info[1:]:
                                                        if abs(word['center_y'] - current_y) <= line_threshold:
                                                            current_line.append(word)
                                                        else:
                                                            lines.append(current_line)
                                                            current_line = [word]
                                                            current_y = word['center_y']
                                                    lines.append(current_line)

                                                for line in lines:
                                                    line.sort(key=lambda x: x['x0'])

                                                for line in lines:
                                                    line_text = " ".join([w['text'] for w in line])
                                                    if merged_text:
                                                        merged_text += "\n" + line_text
                                                    else:
                                                        merged_text = line_text
                                                merged_text = merged_text.strip()

                                            # 添加合并后的文本作为第一列
                                            row_data.append(merged_text)

                                            # 处理剩余列
                                            for col_idx in range(2, len(vertical_lines) - 1):
                                                left = vertical_lines[col_idx]
                                                right = vertical_lines[col_idx + 1]

                                                # 提取该单元格内的文本
                                                cell_bbox = fitz.Rect(left, top, right, bottom)
                                                words = page.get_text("words", clip=cell_bbox)
                                                cell_text = ""
                                                if words:
                                                    word_info = []
                                                    for word in words:
                                                        x0, y0, x1, y1, text, block_no, line_no, word_no = word
                                                        center_y = (y0 + y1) / 2
                                                        word_info.append({
                                                            'text': text,
                                                            'x0': x0,
                                                            'x1': x1,
                                                            'center_y': center_y,
                                                            'height': y1 - y0
                                                        })

                                                    word_info.sort(key=lambda x: x['center_y'])

                                                    lines = []
                                                    if word_info:
                                                        avg_height = sum([w['height'] for w in word_info]) / len(word_info)
                                                        line_threshold = avg_height * 0.5

                                                        current_line = [word_info[0]]
                                                        current_y = word_info[0]['center_y']

                                                        for word in word_info[1:]:
                                                            if abs(word['center_y'] - current_y) <= line_threshold:
                                                                current_line.append(word)
                                                            else:
                                                                lines.append(current_line)
                                                                current_line = [word]
                                                                current_y = word['center_y']
                                                        lines.append(current_line)

                                                    for line in lines:
                                                        line.sort(key=lambda x: x['x0'])

                                                    for line in lines:
                                                        line_text = " ".join([w['text'] for w in line])
                                                        if cell_text:
                                                            cell_text += "\n" + line_text
                                                        else:
                                                            cell_text = line_text
                                                    cell_text = cell_text.strip()

                                                row_data.append(cell_text)
                                    # 第2-7、9-12行的每一行的第2到其后的列合并
                                    # 第14-16行的每一行先合并单元格再提取
                                    elif valid_row_idx in [1, 2, 3, 4, 5, 6, 8, 9, 10, 11, 13, 14, 15]:
                                        if len(vertical_lines) >= 2:
                                            # 第一列保持不变
                                            left1 = vertical_lines[0]
                                            right1 = vertical_lines[1]

                                            # 提取第一列文本
                                            cell_bbox = fitz.Rect(left1, top, right1, bottom)
                                            words = page.get_text("words", clip=cell_bbox)
                                            first_col_text = ""
                                            if words:
                                                word_info = []
                                                for word in words:
                                                    x0, y0, x1, y1, text, block_no, line_no, word_no = word
                                                    center_y = (y0 + y1) / 2
                                                    word_info.append({
                                                        'text': text,
                                                        'x0': x0,
                                                        'x1': x1,
                                                        'center_y': center_y,
                                                        'height': y1 - y0
                                                    })

                                                word_info.sort(key=lambda x: x['center_y'])

                                                lines = []
                                                if word_info:
                                                    avg_height = sum([w['height'] for w in word_info]) / len(word_info)
                                                    line_threshold = avg_height * 0.5

                                                    current_line = [word_info[0]]
                                                    current_y = word_info[0]['center_y']

                                                    for word in word_info[1:]:
                                                        if abs(word['center_y'] - current_y) <= line_threshold:
                                                            current_line.append(word)
                                                        else:
                                                            lines.append(current_line)
                                                            current_line = [word]
                                                            current_y = word['center_y']
                                                    lines.append(current_line)

                                                for line in lines:
                                                    line.sort(key=lambda x: x['x0'])

                                                for line in lines:
                                                    line_text = " ".join([w['text'] for w in line])
                                                    if first_col_text:
                                                        first_col_text += "\n" + line_text
                                                    else:
                                                        first_col_text = line_text
                                                first_col_text = first_col_text.strip()

                                            row_data.append(first_col_text)

                                            # 合并第2到其后的列
                                            if len(vertical_lines) >= 3:
                                                # 对于第14-16行（valid_row_idx=13,14,15），先合并单元格再提取
                                                if valid_row_idx in [13, 14, 15]:
                                                    # 合并第2到其后的所有列为一个大单元格
                                                    left2 = vertical_lines[1]
                                                    right2 = vertical_lines[-1]

                                                    # 提取合并后的单元格文本
                                                    cell_bbox = fitz.Rect(left2, top, right2, bottom)
                                                    words = page.get_text("words", clip=cell_bbox)
                                                    merged_text = ""
                                                    if words:
                                                        word_info = []
                                                        for word in words:
                                                            x0, y0, x1, y1, text, block_no, line_no, word_no = word
                                                            center_y = (y0 + y1) / 2
                                                            word_info.append({
                                                                'text': text,
                                                                'x0': x0,
                                                                'x1': x1,
                                                                'center_y': center_y,
                                                                'height': y1 - y0
                                                            })

                                                        word_info.sort(key=lambda x: x['center_y'])

                                                        lines = []
                                                        if word_info:
                                                            avg_height = sum([w['height'] for w in word_info]) / len(word_info)
                                                            line_threshold = avg_height * 0.5

                                                            current_line = [word_info[0]]
                                                            current_y = word_info[0]['center_y']

                                                            for word in word_info[1:]:
                                                                if abs(word['center_y'] - current_y) <= line_threshold:
                                                                    current_line.append(word)
                                                                else:
                                                                    lines.append(current_line)
                                                                    current_line = [word]
                                                                    current_y = word['center_y']
                                                            lines.append(current_line)

                                                        for line in lines:
                                                            line.sort(key=lambda x: x['x0'])

                                                        for line in lines:
                                                            line_text = " ".join([w['text'] for w in line])
                                                            if merged_text:
                                                                merged_text += "\n" + line_text
                                                            else:
                                                                merged_text = line_text
                                                        merged_text = merged_text.strip()
                                                # 对于第9-12行（valid_row_idx=8,9,10,11），先提取第2列，再合并第3到其后的列
                                                elif valid_row_idx in [8, 9, 10, 11]:
                                                    # 提取第2列文本
                                                    left2 = vertical_lines[1]
                                                    right2 = vertical_lines[2]
                                                    cell_bbox = fitz.Rect(left2, top, right2, bottom)
                                                    words = page.get_text("words", clip=cell_bbox)
                                                    col2_text = ""
                                                    if words:
                                                        word_info = []
                                                        for word in words:
                                                            x0, y0, x1, y1, text, block_no, line_no, word_no = word
                                                            center_y = (y0 + y1) / 2
                                                            word_info.append({
                                                                'text': text,
                                                                'x0': x0,
                                                                'x1': x1,
                                                                'center_y': center_y,
                                                                'height': y1 - y0
                                                            })

                                                        word_info.sort(key=lambda x: x['center_y'])

                                                        lines = []
                                                        if word_info:
                                                            avg_height = sum([w['height'] for w in word_info]) / len(word_info)
                                                            line_threshold = avg_height * 0.5

                                                            current_line = [word_info[0]]
                                                            current_y = word_info[0]['center_y']

                                                            for word in word_info[1:]:
                                                                if abs(word['center_y'] - current_y) <= line_threshold:
                                                                    current_line.append(word)
                                                                else:
                                                                    lines.append(current_line)
                                                                    current_line = [word]
                                                                    current_y = word['center_y']
                                                            lines.append(current_line)

                                                        for line in lines:
                                                            line.sort(key=lambda x: x['x0'])

                                                        for line in lines:
                                                            line_text = " ".join([w['text'] for w in line])
                                                            if col2_text:
                                                                col2_text += "\n" + line_text
                                                            else:
                                                                col2_text = line_text
                                                        col2_text = col2_text.strip()

                                                    # 合并第3到其后的所有列为一个大单元格
                                                    if len(vertical_lines) >= 4:
                                                        left3 = vertical_lines[2]
                                                        right3 = vertical_lines[-1]

                                                        # 提取合并后的单元格文本
                                                        cell_bbox = fitz.Rect(left3, top, right3, bottom)
                                                        words = page.get_text("words", clip=cell_bbox)
                                                        col3_plus_text = ""
                                                        if words:
                                                            word_info = []
                                                            for word in words:
                                                                x0, y0, x1, y1, text, block_no, line_no, word_no = word
                                                                center_y = (y0 + y1) / 2
                                                                word_info.append({
                                                                    'text': text,
                                                                    'x0': x0,
                                                                    'x1': x1,
                                                                    'center_y': center_y,
                                                                    'height': y1 - y0
                                                                })

                                                            word_info.sort(key=lambda x: x['center_y'])

                                                            lines = []
                                                            if word_info:
                                                                avg_height = sum([w['height'] for w in word_info]) / len(word_info)
                                                                line_threshold = avg_height * 0.5

                                                                current_line = [word_info[0]]
                                                                current_y = word_info[0]['center_y']

                                                                for word in word_info[1:]:
                                                                    if abs(word['center_y'] - current_y) <= line_threshold:
                                                                        current_line.append(word)
                                                                    else:
                                                                        lines.append(current_line)
                                                                        current_line = [word]
                                                                        current_y = word['center_y']
                                                                lines.append(current_line)

                                                            for line in lines:
                                                                line.sort(key=lambda x: x['x0'])

                                                            for line in lines:
                                                                line_text = " ".join([w['text'] for w in line])
                                                                if col3_plus_text:
                                                                    col3_plus_text += "\n" + line_text
                                                                else:
                                                                    col3_plus_text = line_text
                                                            col3_plus_text = col3_plus_text.strip()
                                                    else:
                                                        col3_plus_text = ""

                                                    # 合并第2列和第3+列的文本
                                                    merged_text = " "
                                                    if col2_text:
                                                        merged_text = col2_text
                                                    if col3_plus_text:
                                                        if merged_text:
                                                            merged_text += " " + col3_plus_text
                                                        else:
                                                            merged_text = col3_plus_text
                                                else:
                                                    # 对于其他行，遍历第2到其后的每一列，对空单元格替换为空格
                                                    column_texts = []
                                                    for col_idx in range(1, len(vertical_lines) - 1):
                                                        left = vertical_lines[col_idx]
                                                        right = vertical_lines[col_idx + 1]

                                                        # 提取该单元格内的文本
                                                        cell_bbox = fitz.Rect(left, top, right, bottom)
                                                        words = page.get_text("words", clip=cell_bbox)
                                                        cell_text = ""
                                                        if words:
                                                            word_info = []
                                                            for word in words:
                                                                x0, y0, x1, y1, text, block_no, line_no, word_no = word
                                                                center_y = (y0 + y1) / 2
                                                                word_info.append({
                                                                    'text': text,
                                                                    'x0': x0,
                                                                    'x1': x1,
                                                                    'center_y': center_y,
                                                                    'height': y1 - y0
                                                                })

                                                            word_info.sort(key=lambda x: x['center_y'])

                                                            lines = []
                                                            if word_info:
                                                                avg_height = sum([w['height'] for w in word_info]) / len(word_info)
                                                                line_threshold = avg_height * 0.5

                                                                current_line = [word_info[0]]
                                                                current_y = word_info[0]['center_y']

                                                                for word in word_info[1:]:
                                                                    if abs(word['center_y'] - current_y) <= line_threshold:
                                                                        current_line.append(word)
                                                                    else:
                                                                        lines.append(current_line)
                                                                        current_line = [word]
                                                                        current_y = word['center_y']
                                                                lines.append(current_line)

                                                            for line in lines:
                                                                line.sort(key=lambda x: x['x0'])

                                                            for line in lines:
                                                                line_text = " ".join([w['text'] for w in line])
                                                                if cell_text:
                                                                    cell_text += "\n" + line_text
                                                                else:
                                                                    cell_text = line_text
                                                            cell_text = cell_text.strip()

                                                        # 如果单元格文本为空，根据行号决定是否用空格替换
                                                        if not cell_text:
                                                            # 第2-7行（valid_row_idx 1-6）用空格替换
                                                            if valid_row_idx in [1, 2, 3, 4, 5, 6]:
                                                                cell_text = " "
                                                            # 9-12行保持为空
                                                        column_texts.append(cell_text)

                                                    # 合并所有列的文本
                                                    merged_text = " ".join(column_texts)

                                                # 添加合并后的文本作为第二列
                                                row_data.append(merged_text)
                                                # 其余列留空
                                                for _ in range(2, len(vertical_lines) - 1):
                                                    row_data.append("")
                                            else:
                                                row_data.append("")
                                    # 第8行（valid_row_idx=7）不提取
                                    elif valid_row_idx == 7:
                                        pass
                                    else:
                                        # 其他行按照原来的提取规则
                                        for col_idx in range(len(vertical_lines) - 1):
                                            left = vertical_lines[col_idx]
                                            right = vertical_lines[col_idx + 1]

                                            # 提取该单元格内的文本
                                            cell_bbox = fitz.Rect(left, top, right, bottom)
                                            words = page.get_text("words", clip=cell_bbox)
                                            cell_text = ""
                                            if words:
                                                word_info = []
                                                for word in words:
                                                    x0, y0, x1, y1, text, block_no, line_no, word_no = word
                                                    center_y = (y0 + y1) / 2
                                                    word_info.append({
                                                        'text': text,
                                                        'x0': x0,
                                                        'x1': x1,
                                                        'center_y': center_y,
                                                        'height': y1 - y0
                                                    })

                                                word_info.sort(key=lambda x: x['center_y'])

                                                lines = []
                                                if word_info:
                                                    avg_height = sum([w['height'] for w in word_info]) / len(word_info)
                                                    line_threshold = avg_height * 0.5

                                                    current_line = [word_info[0]]
                                                    current_y = word_info[0]['center_y']

                                                    for word in word_info[1:]:
                                                        if abs(word['center_y'] - current_y) <= line_threshold:
                                                            current_line.append(word)
                                                        else:
                                                            lines.append(current_line)
                                                            current_line = [word]
                                                            current_y = word['center_y']
                                                    lines.append(current_line)

                                                for line in lines:
                                                    line.sort(key=lambda x: x['x0'])

                                                for line in lines:
                                                    line_text = " ".join([w['text'] for w in line])
                                                    if cell_text:
                                                        cell_text += "\n" + line_text
                                                    else:
                                                        cell_text = line_text
                                                cell_text = cell_text.strip()

                                            row_data.append(cell_text)
                                # 处理第5个字符串矩形（i=4）的特殊规则
                                elif i == 4:
                                    # 提取该行所有单元格的文本
                                    cells_text = []
                                    for col_idx in range(len(vertical_lines) - 1):
                                        left = vertical_lines[col_idx]
                                        right = vertical_lines[col_idx + 1]

                                        # 提取该单元格内的文本
                                        cell_bbox = fitz.Rect(left, top, right, bottom)
                                        words = page.get_text("words", clip=cell_bbox)
                                        cell_text = ""
                                        if words:
                                            word_info = []
                                            for word in words:
                                                x0, y0, x1, y1, text, block_no, line_no, word_no = word
                                                center_y = (y0 + y1) / 2
                                                word_info.append({
                                                    'text': text,
                                                    'x0': x0,
                                                    'x1': x1,
                                                    'center_y': center_y,
                                                    'height': y1 - y0
                                                })

                                            word_info.sort(key=lambda x: x['center_y'])

                                            lines = []
                                            if word_info:
                                                avg_height = sum([w['height'] for w in word_info]) / len(word_info)
                                                line_threshold = avg_height * 0.5

                                                current_line = [word_info[0]]
                                                current_y = word_info[0]['center_y']

                                                for word in word_info[1:]:
                                                    if abs(word['center_y'] - current_y) <= line_threshold:
                                                        current_line.append(word)
                                                    else:
                                                        lines.append(current_line)
                                                        current_line = [word]
                                                        current_y = word['center_y']
                                                lines.append(current_line)

                                            for line in lines:
                                                line.sort(key=lambda x: x['x0'])

                                            for line in lines:
                                                line_text = " ".join([w['text'] for w in line])
                                                if cell_text:
                                                    cell_text += "\n" + line_text
                                                else:
                                                    cell_text = line_text
                                                cell_text = cell_text.strip()
                                        cells_text.append(cell_text)

                                    # 检查第一个单元格是否为空，而其他至少有一个单元格不为空
                                    if not cells_text[0].strip() and any(cell.strip() for cell in cells_text[1:]):
                                        # 找到属性名为“零序CT安装位置”或“ZCT MOUNTING SITE”的行
                                        for existing_row in table_data:
                                            if len(existing_row) > 0:
                                                row_text = existing_row[0].strip()
                                                if row_text == "零序CT安装位置" or row_text == "ZCT MOUNTING SITE":
                                                    # 将当前行的数据合并到该属性行的数据列后面
                                                    if len(existing_row) > 1:
                                                        # 合并其他单元格的文本
                                                        additional_text = " ".join([cell for cell in cells_text[1:] if cell.strip()])
                                                        if additional_text:
                                                            existing_row[1] += " " + additional_text
                                                    break
                                    else:
                                        # 其他行按照原来的提取规则
                                        for cell_text in cells_text:
                                            row_data.append(cell_text)
                                else:
                                    # 其他字符串矩形，按照原来的提取规则
                                    for col_idx in range(len(vertical_lines) - 1):
                                        left = vertical_lines[col_idx]
                                        right = vertical_lines[col_idx + 1]

                                        # 提取该单元格内的文本
                                        cell_bbox = fitz.Rect(left, top, right, bottom)
                                        words = page.get_text("words", clip=cell_bbox)
                                        cell_text = ""
                                        if words:
                                            word_info = []
                                            for word in words:
                                                x0, y0, x1, y1, text, block_no, line_no, word_no = word
                                                center_y = (y0 + y1) / 2
                                                word_info.append({
                                                    'text': text,
                                                    'x0': x0,
                                                    'x1': x1,
                                                    'center_y': center_y,
                                                    'height': y1 - y0
                                                })

                                            word_info.sort(key=lambda x: x['center_y'])

                                            lines = []
                                            if word_info:
                                                avg_height = sum([w['height'] for w in word_info]) / len(word_info)
                                                line_threshold = avg_height * 0.5

                                                current_line = [word_info[0]]
                                                current_y = word_info[0]['center_y']

                                                for word in word_info[1:]:
                                                    if abs(word['center_y'] - current_y) <= line_threshold:
                                                        current_line.append(word)
                                                    else:
                                                        lines.append(current_line)
                                                        current_line = [word]
                                                        current_y = word['center_y']
                                                lines.append(current_line)

                                            for line in lines:
                                                line.sort(key=lambda x: x['x0'])

                                            for line in lines:
                                                line_text = " ".join([w['text'] for w in line])
                                                if cell_text:
                                                    cell_text += "\n" + line_text
                                                else:
                                                    cell_text = line_text
                                                cell_text = cell_text.strip()

                                        row_data.append(cell_text)

                                # 检查row_data中是否所有列都为空，并且对每一行的元素进行检查，移除字符串元素开头和末尾的空格
                                stripped_row_data = [cell.strip() if isinstance(cell, str) else cell for cell in row_data]
                                if any(cell for cell in stripped_row_data):
                                    table_data.append(stripped_row_data)

                            # 打印表格数据
                            print(f"字符串 {i + 1} 表格数据:")
                            for row in table_data:
                                print(row)

                            all_tables_data.append(table_data)

                            # 绘制表格网格（只绘制高度足够的行）
                            shape = page.new_shape()
                            # 收集有效的水平分割线（忽略高度很小的行）
                            valid_horizontal_lines = []
                            min_row_height = 2  # 最小行高阈值
                            for row_idx in range(len(horizontal_lines)):
                                if row_idx == 0:
                                    # 保留第一条线
                                    valid_horizontal_lines.append(horizontal_lines[row_idx])
                                else:
                                    # 检查与前一条线的高度
                                    prev_y = horizontal_lines[row_idx - 1]
                                    current_y = horizontal_lines[row_idx]
                                    if current_y - prev_y >= min_row_height:
                                        valid_horizontal_lines.append(current_y)

                            # 绘制水平分割线
                            for y in valid_horizontal_lines:
                                shape.draw_line(fitz.Point(vertical_lines[0], y), fitz.Point(vertical_lines[-1], y))
                            # 绘制垂直分割线
                            for x in vertical_lines:
                                if valid_horizontal_lines:
                                    shape.draw_line(fitz.Point(x, valid_horizontal_lines[0]), fitz.Point(x, valid_horizontal_lines[-1]))
                            shape.finish(color=(0, 1, 1), width=1)
                            shape.commit()

                            # 在标注图像中绘制这个区域
                            shape = page.new_shape()
                            shape.draw_rect(fitz.Rect(best_region[0], best_region[1], best_region[2], best_region[3]))
                            shape.finish(color=(1, 0, 1), fill=None, width=2, dashes=[2, 2])
                            shape.commit()
                        else:
                            string_regions.append(None)
                            print(f"字符串 {i + 1} 未找到对应的矩形区域")
            else:
                print("未检测到字符串，生成页面边界标注...")
                # 绘制页面边界
                shape = page.new_shape()
                shape.draw_rect(fitz.Rect(0, 0, page.rect.width, page.rect.height))
                shape.finish(color=(1, 0, 0), fill=None, width=2)
                shape.commit()

            # 将页面转换为图像
            mat = fitz.Matrix(2, 2)  # 放大2倍以提高清晰度
            pix = page.get_pixmap(matrix=mat)

            # 保存图像
            try:
                pix.save(output_image_path)
                print(f"矩形标注图像已保存到: {output_image_path}\n")
            except Exception as e:
                print(f"保存标注图像时出错: {e}")

        doc.close()

        return all_tables_data

    except Exception as e:
        print(f"提取有效矩形时出错: {e}")
        return []


def extract_tables_from_pdf(pdf_path, search_terms, sheet_1_7_page_number, table_output_folder='crop/extracted_tables', image_output_folder='crop/extracted_images'):
    all_tables_data = {}

    if not os.path.exists(table_output_folder):
        os.makedirs(table_output_folder)
    if not os.path.exists(image_output_folder):
        os.makedirs(image_output_folder)

    with pdfplumber.open(pdf_path) as pdf:
        for page_number, page in enumerate(pdf.pages):
            if page_number in sheet_1_7_page_number:
                tables_data = extract_tables_method_1(page, page_number, search_terms=search_terms, table_output_folder=table_output_folder, image_output_folder=image_output_folder, pdf_path=pdf_path)
                if tables_data:
                    all_tables_data[page_number] = tables_data

    return all_tables_data


def extract_tables_method_1(page, page_number, search_terms, table_output_folder, image_output_folder, pdf_path=None):
    print(f"提取页面中的有效矩形区域...")

    rectangle_annotation_path = None
    if pdf_path and image_output_folder:
        pdf_filename = os.path.splitext(os.path.basename(pdf_path))[0]
        rectangle_annotation_path = os.path.join(image_output_folder, f'{pdf_filename}_page_{page_number + 1}_annotation.png')

    if pdf_path:
        tables_data = extract_valid_rectangles(pdf_path, page_number, search_terms, output_image_path=rectangle_annotation_path)
        return tables_data
    else:
        print("未提供PDF路径，使用整个页面作为区域")
        return []


def load_config_file(config_file):
    """
    读取配置文件，返回中英文属性名称映射、EPLAN报表序号映射和拼接符映射

    参数:
        config_file: 配置文件路径

    返回:
        tuple: (name_mapping, eplan_mapping, separator_mapping)
            name_mapping: 中英文属性名称映射
            eplan_mapping: 属性名称到EPLAN报表序号的映射
            separator_mapping: 属性名称到拼接符的映射 (separator1, separator2, separator3)
    """

    try:
        wb = load_workbook(config_file)
        ws = wb.active

        name_mapping = {}
        eplan_mapping = {}
        separator_mapping = {}

        # 假设第一行是列名，从第二行开始读取数据
        for row in range(2, ws.max_row + 1):
            chinese = ws.cell(row=row, column=1).value
            english = ws.cell(row=row, column=2).value
            eplan = ws.cell(row=row, column=3).value
            separator1 = ws.cell(row=row, column=4).value
            separator2 = ws.cell(row=row, column=5).value
            separator3 = ws.cell(row=row, column=6).value

            if chinese and english:
                name_mapping[english] = chinese
            if chinese and eplan:
                eplan_mapping[chinese] = eplan
            if chinese:
                separator_mapping[chinese] = (separator1, separator2, separator3)

        return name_mapping, eplan_mapping, separator_mapping
    except Exception as e:
        print(f"读取配置文件出错: {e}")
        return {}, {}, {}


def create_gui():
    """
    创建GUI界面
    """
    root = tk.Tk()
    root.title("DrawMaster_DS_V1.0_20260518—DataSheet参数提取及对比工具")
    root.geometry("1000x700")

    try:
        root.iconbitmap("logo.ico")
    except Exception:
        pass  # 如果图标文件不存在，忽略错误

    # 窗口最大化
    root.state('zoomed')

    # 窗口居中显示
    root.update_idletasks()
    width = root.winfo_width()
    height = root.winfo_height()
    x = (root.winfo_screenwidth() // 2) - (width // 2)
    y = (root.winfo_screenheight() // 2) - (height // 2)
    root.geometry(f"{width}x{height}+{x}+{y}")

    # 设置字体
    font = ("ABBvoice CNSG", 10)
    big_font = ("ABBvoice CNSG", 12)

    # 设置Treeview字体和行高
    style = ttk.Style()
    # 普通行样式（默认高度）
    style.configure("Treeview", font=font, rowheight=25)
    # 两倍高度行样式
    style.configure("Double.Treeview", font=font, rowheight=50)
    # 四倍高度行样式
    style.configure("Quadruple.Treeview", font=font, rowheight=100)
    style.configure("Treeview.Heading", font=font)

    # 顶部文件选择区域
    top_frame = tk.Frame(root, padx=10, pady=10)
    top_frame.pack(fill=tk.X)

    # 文件路径Entry
    tk.Label(top_frame, text="PDF文件:", font=big_font).pack(side=tk.LEFT, padx=5)

    pdf_path_var = tk.StringVar()
    pdf_entry = tk.Entry(top_frame, textvariable=pdf_path_var, width=80, font=big_font)
    pdf_entry.pack(side=tk.LEFT, padx=5)

    # 选择文件按钮
    def select_pdf():
        # 清空文件路径Entry
        pdf_path_var.set("")

        # 清空识别属性Canvas
        if hasattr(canvas_inner_frame, 'winfo_children'):
            for widget in canvas_inner_frame.winfo_children():
                widget.destroy()
            # 重置canvas_inner_frame的属性，确保不再保留对已销毁组件的引用
            # 保留必要的属性如tk、master等
            attributes_to_preserve = ['tk', 'master', 'name', '_w', '_name', 'children']
            for attr in list(canvas_inner_frame.__dict__.keys()):
                if not attr.startswith('__') and attr not in attributes_to_preserve:
                    delattr(canvas_inner_frame, attr)
            # 更新Canvas滚动区域
            canvas_inner_frame.update_idletasks()
            recognition_canvas.config(scrollregion=recognition_canvas.bbox("all"))

        # 清空识别结果Text
        recognition_text.config(state=tk.NORMAL)
        recognition_text.delete(1.0, tk.END)
        recognition_text.insert(tk.END, "识别结果将显示在这里\n")
        recognition_text.config(state=tk.DISABLED)

        # 清空对比结果Text
        compare_text.config(state=tk.NORMAL)
        compare_text.delete(1.0, tk.END)
        compare_text.insert(tk.END, "对比结果将显示在这里\n")
        compare_text.config(state=tk.DISABLED)

        # 清空对比值缓存
        compare_values_cache.clear()

        # 打开文件选择对话框
        file_path = filedialog.askopenfilename(
            title="选择PDF文件",
            filetypes=[("PDF文件", "*.pdf"), ("所有文件", "*")]
        )
        if file_path:
            pdf_path_var.set(file_path)
            # 启用识别按钮
            recognize_button.config(state=tk.NORMAL)
            # 禁用对比和导出按钮
            compare_button.config(state=tk.DISABLED)
            export_button.config(state=tk.DISABLED)

    tk.Button(top_frame, text="选择图纸", command=select_pdf, font=big_font, width=8).pack(side=tk.LEFT, padx=5)

    # 识别按钮
    def start_recognition():
        pdf_path = pdf_path_var.get()
        if not pdf_path or not os.path.exists(pdf_path):
            messagebox.showerror("错误", "请选择有效的PDF文件")
            return

        # 清空两个Text区域
        # 清空左边识别结果
        recognition_text.config(state=tk.NORMAL)
        recognition_text.delete(1.0, tk.END)
        recognition_text.insert(tk.END, f"{'-' * 60}\n")
        recognition_text.insert(tk.END, f"开始处理: {pdf_path}\n")
        recognition_text.insert(tk.END, f"{'-' * 60}\n")
        recognition_text.config(state=tk.DISABLED)

        # 清空左边Canvas
        for widget in canvas_inner_frame.winfo_children():
            widget.destroy()

        # 清空右边对比结果
        compare_text.config(state=tk.NORMAL)
        compare_text.delete(1.0, tk.END)
        compare_text.config(state=tk.DISABLED)

        # 清空对比值缓存
        compare_values_cache.clear()

        # 清空页面数据
        all_pages_data.clear()

        # 重定向print输出到左边Text
        import sys
        class TextRedirector:
            def __init__(self, widget):
                self.widget = widget
                self.console = sys.stdout  # 保存原始的stdout

            def write(self, text):
                # 输出到控制台
                # self.console.write(text)
                # 输出到Text
                self.widget.config(state=tk.NORMAL)
                self.widget.insert(tk.END, text)
                self.widget.see(tk.END)
                self.widget.config(state=tk.DISABLED)

            def flush(self):
                self.console.flush()

        old_stdout = sys.stdout
        sys.stdout = TextRedirector(recognition_text)

        try:
            # 执行原有的处理逻辑
            search_terms = [
                '开关站基本参数', 'SWITCHBOARD BASIC DATA',  # 开关站基本参数
                '铭牌', 'LABELS',  # 铭牌
                '辅助电源', 'AUX POWER SUPPLY',  # 辅助电源
                '二次导线参数', 'SECONDARY WIRES PARAMETERS',  # 二次导线参数
                '其他', 'OTHER',  # 其他
                '检验', 'INSPECTION',  # 检验
                '包装和发运', 'PACKING DETAILS & SHIPPING MARKS'  # 包装和发运
            ]

            sheet_1_7_page_number = []
            page_number_mapping = {}

            with open(pdf_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                outlines = pdf_reader.outline

                for index, item in enumerate(outlines[3]):
                    title = item.get("/Title", "")
                    if '&DS' in title and ('技术参数表' in title or 'TECHNICAL DATA SHEET' in title):
                        # print(index, item)
                        sheet_1_7_page_number.append(index)

                        switchgear_num = extract_switchgear_number(title)
                        if switchgear_num:
                            page_number_mapping[index] = switchgear_num
                            # print(f"  提取到站号: {switchgear_num}")

            if sheet_1_7_page_number:
                project_number = find_project_number(pdf_path, sheet_1_7_page_number[0])
                if not project_number:
                    # 未找到项目号，使用PDF文件名作为项目号
                    pdf_filename = os.path.splitext(os.path.basename(pdf_path))[0]
                    project_number = pdf_filename
                    print(f"未找到项目号，使用PDF文件名作为项目号: {project_number}")
                else:
                    print(f"\n找到的项目号: {project_number}")

                table_output_folder = 'crop/extracted_tables'
                image_output_folder = 'crop/extracted_images'

                # 确保目录存在
                if not os.path.exists(table_output_folder):
                    os.makedirs(table_output_folder)
                if not os.path.exists(image_output_folder):
                    os.makedirs(image_output_folder)

                all_tables_data = extract_tables_from_pdf(pdf_path, search_terms, sheet_1_7_page_number, table_output_folder, image_output_folder)

                wb = Workbook()

                for page_num, tables_data in all_tables_data.items():
                    switchgear_num = page_number_mapping.get(page_num, f"Page_{page_num}")

                    if switchgear_num:
                        ws_new = wb.create_sheet(title=switchgear_num)

                        row_offset = 1
                        for table_data in tables_data:
                            for row in table_data:
                                for col_idx, cell_value in enumerate(row, 1):
                                    ws_new.cell(row=row_offset, column=col_idx, value=cell_value)
                                row_offset += 1
                            row_offset += 1

                # 删除默认的Sheet1
                if 'Sheet' in wb.sheetnames:
                    wb.remove(wb['Sheet'])

                pdf_filename = os.path.splitext(os.path.basename(pdf_path))[0]
                output_filename = os.path.join(table_output_folder, f'{pdf_filename}.xlsx')
                wb.save(output_filename)
                print(f"Excel文件已保存: {output_filename}")

                # 读取配置文件
                # 默认配置文件路径：主程序同级的attributes_mappinglist.xlsx
                config_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), "attributes_mappinglist.xlsx")
                name_mapping = {}
                eplan_mapping = {}
                separator_mapping = {}
                if os.path.exists(config_file):
                    name_mapping, eplan_mapping, separator_mapping = load_config_file(config_file)
                    print(f"加载配置文件成功，映射数量: {len(name_mapping)}, EPLAN映射数量: {len(eplan_mapping)}, 拼接符映射数量: {len(separator_mapping)}")
                else:
                    print(f"配置文件不存在: {config_file}")

                # 存储所有识别的属性组（按页面分组）
                attribute_groups = []
                for page_num, tables_data in all_tables_data.items():
                    page_trees = []
                    max_table_idx = len(tables_data) - 1  # 获取最大table_idx
                    for table_idx, table_data in enumerate(tables_data):
                        if table_data:
                            parent_title = search_terms[table_idx * 2]

                            if not parent_title:
                                parent_title = f"表格 {table_idx + 1}"
                            # 子节点数据
                            children = []
                            for row in table_data[1:]:
                                # 移除包含特定字符串的行
                                row_str = "".join(str(item) for item in row)
                                if "注：" in row_str or "SWITCHGEAR层级" in row_str or "项目层级" in row_str:
                                    continue

                                if table_idx == max_table_idx and len(row) == 1 and row[0]:
                                    # 特殊处理：当table_idx为最大值时，len(row) == 1且row[0]存在
                                    attr_name = "包装发运"
                                    attr_value = row[0]
                                elif len(row) >= 2:
                                    # 将row[1]及其后的元素用\n拼接
                                    attr_name = row[0]
                                    value = "\n".join(str(item) for item in row[1:] if item)
                                    attr_value = value
                                elif len(row) == 1 and row[0]:
                                    attr_name = row[0]
                                    attr_value = ""
                                else:
                                    continue

                                # 处理属性名称，根据配置文件统一替换为纯中文
                                if name_mapping:
                                    # 检查属性名称是否在映射中
                                    if attr_name in name_mapping:
                                        attr_name = name_mapping[attr_name]
                                    else:
                                        # 检查是否是英文或中英文混合
                                        # 简单判断：如果包含字母，则认为是英文或中英文混合
                                        if any(c.isalpha() for c in str(attr_name)):
                                            # 尝试查找配置文件中的英文映射
                                            for english, chinese in name_mapping.items():
                                                if english in str(attr_name):
                                                    attr_name = chinese
                                                    break

                                children.append((attr_name, attr_value))
                            if children:
                                page_trees.append((parent_title, children))
                    if page_trees:
                        # 为每个属性组添加标题（使用站号）
                        switchgear_num = page_number_mapping.get(page_num, f"Page_{page_num}")
                        page_title = f"{switchgear_num}"
                        attribute_groups.append((page_title, page_trees))

                # 分页显示（按页面分组）
                total_pages = len(attribute_groups)
                current_page = 0  # 从0开始索引

                # 收集所有页面的数据
                collected_pages_data = []
                for page_index in range(total_pages):
                    page_title, page_trees = attribute_groups[page_index]
                    page_data = {
                        'title': page_title,
                        'attributes': []
                    }

                    # 遍历当前页面的所有表格
                    for parent_title, children in page_trees:
                        # 只显示对应search_terms中的0,2,4,6,8,10,12位元素
                        if parent_title in search_terms and search_terms.index(parent_title) % 2 == 0:
                            # 遍历当前表格的所有属性
                            for attr_name, attr_value in children:
                                # 从缓存中获取对比值
                                old_compare_value = ""
                                if attr_name in compare_values_cache:
                                    old_compare_value = compare_values_cache[attr_name]

                                # 存储属性数据
                                page_data['attributes'].append({
                                    'name': attr_name,
                                    'recognized_value': attr_value,
                                    'compare_value': old_compare_value
                                })
                    collected_pages_data.append(page_data)

                # 清空并填充all_pages_data
                all_pages_data.clear()
                all_pages_data.extend(collected_pages_data)

                def show_page(page_index):
                    # 清空Canvas
                    for widget in canvas_inner_frame.winfo_children():
                        widget.destroy()

                    if 0 <= page_index < total_pages:
                        # 获取当前属性组
                        page_title, page_trees = attribute_groups[page_index]

                        # 存储当前页面的站号
                        canvas_inner_frame.page_title = page_title

                        # 添加页面标题
                        page_title_label = tk.Label(canvas_inner_frame, text=page_title, font=(font[0], font[1] + 2, 'bold'), bg='white', anchor='w')
                        page_title_label.pack(fill=tk.X, pady=(10, 5))

                        # 存储所有compare_value_text组件的引用
                        if not hasattr(canvas_inner_frame, 'compare_value_texts'):
                            canvas_inner_frame.compare_value_texts = {}
                        else:
                            canvas_inner_frame.compare_value_texts.clear()

                        # 存储所有recognized_value_text组件的引用
                        if not hasattr(canvas_inner_frame, 'recognized_value_texts'):
                            canvas_inner_frame.recognized_value_texts = {}
                        else:
                            canvas_inner_frame.recognized_value_texts.clear()

                        # 显示当前页的所有表格（只显示对应search_terms中的0,2,4,6,8,10,12位元素）
                        for parent_title, children in page_trees:
                            # 只显示对应search_terms中的0,2,4,6,8,10,12位元素
                            # 检查parent_title是否在search_terms的偶数索引位置
                            if parent_title in search_terms and search_terms.index(parent_title) % 2 == 0:
                                # 添加表格标题
                                table_title_label = tk.Label(canvas_inner_frame, text=parent_title, font=(font[0], font[1], 'bold'), bg='white', anchor='w')
                                table_title_label.pack(fill=tk.X, pady=(10, 5))

                                # 添加表格内容
                                for attr_name, attr_value in children:
                                    # 创建属性行框架
                                    attr_frame = tk.Frame(canvas_inner_frame, bg='white')
                                    attr_frame.pack(fill=tk.X, pady=2)

                                    # 计算高度（根据换行符数量）
                                    attr_name_lines = str(attr_name).count('\n') + 1
                                    attr_value_lines = str(attr_value).count('\n') + 1
                                    max_lines = max(attr_name_lines, attr_value_lines)

                                    # 属性名称
                                    attr_name_text = tk.Text(attr_frame, font=font, width=50, height=max_lines, wrap=tk.WORD,
                                                             bg='#f8f9fa', fg='#333333', borderwidth=0, relief='solid')
                                    attr_name_text.pack(side=tk.LEFT, padx=5)
                                    attr_name_text.insert(tk.END, attr_name)
                                    attr_name_text.config(state=tk.DISABLED)

                                    # 识别的值
                                    recognized_value_text = tk.Text(attr_frame, font=font, width=60, height=max_lines, wrap=tk.WORD,
                                                                    bg='#ffffff', fg='#333333', borderwidth=1, relief='solid')
                                    recognized_value_text.pack(side=tk.LEFT, padx=5)
                                    recognized_value_text.insert(tk.END, attr_value)
                                    recognized_value_text.config(state=tk.DISABLED)

                                    # 对比表格读取的值
                                    # 这里需要从对比表格中获取对应的值，暂时显示为空
                                    compare_value_text = tk.Text(attr_frame, font=font, width=60, height=max_lines, wrap=tk.WORD,
                                                                 bg='#f0f8ff', fg='#333333', borderwidth=1, relief='solid')
                                    compare_value_text.pack(side=tk.LEFT, padx=5)

                                    # 保留之前的对比值
                                    old_compare_value = ""
                                    # 从缓存中获取对比值
                                    if attr_name in compare_values_cache:
                                        old_compare_value = compare_values_cache[attr_name]
                                    compare_value_text.insert(tk.END, old_compare_value)
                                    compare_value_text.config(state=tk.DISABLED)

                                    # 检查是否需要标红
                                    recognized_value = attr_value.strip()
                                    compare_value = old_compare_value.strip()
                                    if recognized_value != compare_value:
                                        # 标红对比值
                                        compare_value_text.config(state=tk.NORMAL)
                                        compare_value_text.delete(1.0, tk.END)
                                        compare_value_text.tag_configure("red", background="red", foreground="white")
                                        compare_value_text.insert(tk.END, old_compare_value if old_compare_value else " ", "red")
                                        compare_value_text.config(state=tk.DISABLED)

                                    # 存储Text组件的引用，以便后续更新
                                    attr_name_text.compare_value_text = compare_value_text

                                    # 存储所有compare_value_text组件的引用
                                    canvas_inner_frame.compare_value_texts[attr_name] = compare_value_text

                                    # 存储所有recognized_value_text组件的引用
                                    canvas_inner_frame.recognized_value_texts[attr_name] = recognized_value_text

                        # 更新Canvas滚动区域
                        canvas_inner_frame.update_idletasks()
                        recognition_canvas.config(scrollregion=recognition_canvas.bbox("all"))

                        # 更新分页标签
                        page_label.config(text=f"第 {page_index + 1} 页，共 {total_pages} 页")

                        # # 显示当前页的对比结果
                        show_current_page_comparison()
                    else:
                        # 空页面
                        empty_label = tk.Label(canvas_inner_frame, text="无数据", font=font, bg='white')
                        empty_label.pack(pady=20)
                        page_label.config(text=f"第 1 页，共 1 页")

                def prev_page():
                    nonlocal current_page
                    if current_page > 0:
                        current_page -= 1
                        show_page(current_page)

                def next_page():
                    nonlocal current_page
                    if current_page < total_pages - 1:
                        current_page += 1
                        show_page(current_page)

                # 绑定按钮事件
                prev_button.config(command=prev_page)
                next_button.config(command=next_page)

                # 显示第一页
                show_page(current_page)
            else:
                print("未找到技术参数表页面")

        except Exception as e:
            print(f"错误: {e}")
        finally:
            sys.stdout = old_stdout
            # 启用对比按钮
            compare_button.config(state=tk.NORMAL)
            # 禁用导出按钮
            export_button.config(state=tk.DISABLED)

    # 识别按钮
    recognize_button = tk.Button(top_frame, text="识别图纸取值", command=start_recognition, font=big_font, width=12, bg="#4CAF50", fg="white")
    recognize_button.pack(side=tk.LEFT, padx=5)

    # 对比按钮
    def start_compare():
        pdf_path = pdf_path_var.get()
        if not pdf_path or not os.path.exists(pdf_path):
            messagebox.showerror("错误", "请先选择并识别PDF文件")
            return

        # 清空右边对比结果
        compare_text.config(state=tk.NORMAL)
        compare_text.delete(1.0, tk.END)
        compare_text.config(state=tk.DISABLED)

        # 选择要对比的Excel文件
        compare_file = filedialog.askopenfilename(
            title="选择要对比的Excel文件",
            filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*")]
        )

        if compare_file:
            # 获取当前生成的Excel文件路径
            table_output_folder = 'crop/extracted_tables'
            pdf_filename = os.path.splitext(os.path.basename(pdf_path))[0]
            current_excel = os.path.join(table_output_folder, f'{pdf_filename}.xlsx')

            if os.path.exists(current_excel):
                # 读取对比文件并填充compare_value_text
                try:
                    wb = load_workbook(compare_file)

                    # 检查canvas_inner_frame是否存在且有compare_value_texts属性
                    if hasattr(canvas_inner_frame, 'compare_value_texts') and hasattr(canvas_inner_frame, 'recognized_value_texts'):
                        compare_values = {}

                        # 只读取sheet名为Z8_xlsx的表单
                        if 'Z8_xlsx' in wb.sheetnames:
                            ws = wb['Z8_xlsx']
                            max_row = ws.max_row
                            max_col = ws.max_column

                            # 读取所有行数据，存储为字典，键为A列的值，值为B列的值
                            compare_data = {}
                            # 存储站号对应的主母线电流和低压室高度
                            station_data = {}

                            # 查找标题行为行号、额定主母线电流、低压室高度的行
                            header_row = -1
                            for row in range(1, max_row + 1):
                                if max_col >= 3:
                                    cell1 = ws.cell(row=row, column=1).value
                                    cell2 = ws.cell(row=row, column=2).value
                                    cell3 = ws.cell(row=row, column=3).value

                                    # 检查是否为目标标题行
                                    if (cell1 and '行号' in str(cell1)) and \
                                            (cell2 and '额定主母线电流' in str(cell2)) and \
                                            (cell3 and '低压室高度' in str(cell3)):
                                        header_row = row
                                        break

                            # 如果找到标题行，读取后续行的站号、主母线电流和低压室高度
                            if header_row > 0:
                                for row in range(header_row + 1, max_row + 1):
                                    if max_col >= 3:
                                        station = ws.cell(row=row, column=1).value
                                        bus_current = ws.cell(row=row, column=2).value
                                        height = ws.cell(row=row, column=3).value

                                        if station:
                                            station_str = str(station).strip()
                                            # 如果站号不存在，或者存在但数据为空或不同，则更新
                                            if station_str not in station_data:
                                                # 站号不存在，创建新条目
                                                station_data[station_str] = {
                                                    '主母线电流': bus_current,
                                                    '低压室高度': height
                                                }
                                            else:
                                                # 站号存在，只更新非空且不同的值
                                                if bus_current is not None and bus_current != station_data[station_str].get('主母线电流'):
                                                    station_data[station_str]['主母线电流'] = bus_current
                                                if height is not None and height != station_data[station_str].get('低压室高度'):
                                                    station_data[station_str]['低压室高度'] = height
                            # 读取所有行数据，存储为字典，键为A列的值，值为B列的值
                            for row in range(1, max_row + 1):
                                if max_col >= 2:
                                    key = ws.cell(row=row, column=1).value
                                    value = ws.cell(row=row, column=2).value
                                    if key:
                                        compare_data[key] = value

                            # 读取配置文件获取EPLAN报表序号映射和拼接符映射
                            config_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), "attributes_mappinglist.xlsx")
                            name_mapping, eplan_mapping, separator_mapping = load_config_file(config_file)

                            # 获取当前页面的站号（如果有）
                            current_station = ""
                            if hasattr(canvas_inner_frame, 'page_title'):
                                current_station = canvas_inner_frame.page_title

                            # 遍历所有属性名称，根据EPLAN报表序号查找对比值
                            for attr_name in canvas_inner_frame.compare_value_texts.keys():
                                # 处理主母线电流和低压室高度的特殊情况
                                # 处理属性名称中带有冒号的情况
                                normalized_attr_name = attr_name.rstrip(':')
                                if normalized_attr_name in ['主母线电流', '低压室高度'] and current_station in station_data:
                                    compare_values[attr_name] = str(station_data[current_station][normalized_attr_name]) if station_data[current_station][normalized_attr_name] else " "
                                elif attr_name in eplan_mapping:
                                    # 获取EPLAN报表序号
                                    eplan_codes = eplan_mapping[attr_name]
                                    if eplan_codes:
                                        # 分割多个EPLAN报表序号
                                        codes = [code.strip() for code in str(eplan_codes).split(',')]
                                        values = []
                                        # 查找每个EPLAN报表序号对应的对比值
                                        for code in codes:
                                            # 查找以code开头的A列值
                                            for key, value in compare_data.items():
                                                if key and str(key).startswith(code):
                                                    values.append(str(value) if value else " ")
                                                    break
                                        # 拼接多个值
                                        if values:
                                            # 获取拼接符
                                            separator1 = "\n"
                                            separator2 = "\n"
                                            separator3 = "\n"
                                            if attr_name in separator_mapping:
                                                sep1, sep2, sep3 = separator_mapping[attr_name]
                                                if sep1:
                                                    separator1 = sep1
                                                if sep2:
                                                    separator2 = sep2
                                                if sep3:
                                                    separator3 = sep3

                                            # 根据,的数量选择拼接符
                                            if len(codes) == 2:
                                                # 有一个,，使用拼接符1
                                                compare_values[attr_name] = separator1.join(values)
                                            elif len(codes) == 3:
                                                # 有两个,，使用拼接符1和拼接符2
                                                if len(values) >= 2:
                                                    compare_values[attr_name] = separator1.join(values[:2]) + separator2 + values[2] if len(values) > 2 else separator1.join(values)
                                            elif len(codes) == 4:
                                                # 有三个,，使用拼接符1、拼接符2和拼接符3
                                                if len(values) >= 3:
                                                    compare_values[attr_name] = separator1.join(values[:2]) + separator2 + values[2] + separator3 + values[3] if len(values) > 3 else separator1.join(values[:2]) + separator2 + values[2]
                                            else:
                                                # 其他情况
                                                compare_values[attr_name] = values[0]

                                # 对于内部燃弧等级，在开头增加'IAC '
                                if attr_name == "内部燃弧等级" and attr_name in compare_values:
                                    compare_values[attr_name] = "IAC " + compare_values[attr_name]
                        else:
                            print("对比文件中没有Z8_xlsx表单")

                        # 填充对比值到compare_value_text
                        # 只处理当前存在的Text组件
                        current_attrs = set(canvas_inner_frame.compare_value_texts.keys())
                        for attr_name, compare_value in compare_values.items():
                            # 存储到缓存中
                            compare_values_cache[attr_name] = compare_value

                            # 更新all_pages_data中的对比值
                            for page_data in all_pages_data:
                                for attr in page_data['attributes']:
                                    if attr['name'] == attr_name:
                                        attr['compare_value'] = compare_value
                                        break

                            if attr_name in current_attrs:
                                text_widget = canvas_inner_frame.compare_value_texts[attr_name]
                                try:
                                    # 检查Text组件是否有效
                                    text_widget.winfo_exists()
                                    text_widget.config(state=tk.NORMAL)
                                    text_widget.delete(1.0, tk.END)
                                    text_widget.insert(tk.END, compare_value if compare_value else "")
                                    text_widget.config(state=tk.DISABLED)
                                except Exception as e:
                                    # 不再打印错误信息，避免干扰用户
                                    continue

                        # 显示当前页的对比结果
                        show_current_page_comparison()

                except Exception as e:
                    print(f"填充对比值出错: {e}")
                    compare_text.config(state=tk.NORMAL)
                    compare_text.insert(tk.END, f"对比出错: {e}\n")
                    compare_text.config(state=tk.DISABLED)

            else:
                messagebox.showerror("错误", "当前PDF尚未生成Excel文件")
        # 启用导出按钮
        export_button.config(state=tk.NORMAL)

        # 强制刷新主窗口，确保canvas内容绘制完成
        root.update_idletasks()
        root.update()

        # 延迟显示统计信息窗口，确保canvas绘制完成
        root.after(100, show_summary_statistics)

    def show_summary_statistics():
        """
        显示统计信息窗口，展示属性和统计情况（Summary的第1,2列）
        """
        if not all_pages_data:
            return

        # 创建Toplevel窗口
        stats_window = tk.Toplevel(root)
        stats_window.title("各站属性取值对比")
        stats_window.geometry("1000x800")
        stats_window.transient(root)  # 设置为父窗口的临时窗口
        stats_window.grab_set()  # 模态窗口
        stats_window.iconbitmap("logo.ico")
        # 窗口居中显示
        stats_window.update_idletasks()
        screen_width = stats_window.winfo_screenwidth()
        screen_height = stats_window.winfo_screenheight()
        x = (screen_width - stats_window.winfo_width()) // 2
        y = (screen_height - stats_window.winfo_height()) // 2
        stats_window.geometry(f"+{x}+{y}")

        # 创建主框架
        main_frame = tk.Frame(stats_window, padx=10, pady=10)
        main_frame.pack(fill=tk.BOTH, expand=True)

        # # 标题
        # title_label = tk.Label(main_frame, text="各站图纸取值对比", font=(font[0], font[1], 'bold'))
        # title_label.pack(pady=(0, 10))

        # 创建Canvas和滚动条
        canvas_frame = tk.Frame(main_frame)
        canvas_frame.pack(fill=tk.BOTH, expand=True)

        canvas = tk.Canvas(canvas_frame, bg='white')
        scrollbar_y = tk.Scrollbar(canvas_frame, orient=tk.VERTICAL, command=canvas.yview)
        scrollbar_x = tk.Scrollbar(canvas_frame, orient=tk.HORIZONTAL, command=canvas.xview)

        canvas.configure(yscrollcommand=scrollbar_y.set, xscrollcommand=scrollbar_x.set)

        scrollbar_y.pack(side=tk.RIGHT, fill=tk.Y)
        scrollbar_x.pack(side=tk.BOTTOM, fill=tk.X)
        canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        # 创建内部框架
        inner_frame = tk.Frame(canvas, bg='white')
        canvas_window = canvas.create_window((0, 0), window=inner_frame, anchor='nw')

        # 表头
        header_frame = tk.Frame(inner_frame, bg='#E0E0E0')
        header_frame.pack(fill=tk.X, pady=(0, 5))

        attr_header = tk.Label(header_frame, text="属性", font=(font[0], font[1], 'bold'),
                               bg='#E0E0E0', width=30, anchor='w')
        attr_header.pack(side=tk.LEFT, padx=5)

        stats_header = tk.Label(header_frame, text="各站图纸取值", font=(font[0], font[1], 'bold'),
                                bg='#E0E0E0', width=50, anchor='w')
        stats_header.pack(side=tk.LEFT, padx=5)

        # 获取所有属性
        all_attributes = []
        attr_set = set()
        for page_data in all_pages_data:
            for attr in page_data['attributes']:
                if attr['name'] not in attr_set:
                    attr_set.add(attr['name'])
                    all_attributes.append(attr['name'])

        station_names = [page_data['title'] for page_data in all_pages_data]

        # 填充数据
        for attr_name in all_attributes:
            values_by_station = {}
            for page_data in all_pages_data:
                station_name = page_data['title']
                value = ""
                for attr in page_data['attributes']:
                    if attr['name'] == attr_name:
                        value = attr['recognized_value']
                        break
                values_by_station[station_name] = value if value.strip() else " "

            # 生成统计信息
            stats, value_count = generate_statistics(station_names, values_by_station)

            # 创建行框架
            row_frame = tk.Frame(inner_frame, bg='white')
            row_frame.pack(fill=tk.X, pady=2)

            # 计算行高
            stats_lines = stats.count('\n') + 1
            row_height = max(1, stats_lines)

            # 确定背景色和字体颜色（有多种值时标红）
            has_multiple_values = value_count > 1
            attr_bg = '#ff0000' if has_multiple_values else '#f8f9fa'
            attr_fg = '#ffffff' if has_multiple_values else '#333333'
            stats_bg = '#ff0000' if has_multiple_values else '#ffffff'
            stats_fg = '#ffffff' if has_multiple_values else '#333333'

            # 属性名称
            attr_text = tk.Text(row_frame, font=font, width=30, height=row_height,
                                wrap=tk.WORD, bg=attr_bg, fg=attr_fg, borderwidth=1, relief='solid')
            attr_text.pack(side=tk.LEFT, padx=5)
            attr_text.insert(tk.END, attr_name)
            attr_text.config(state=tk.DISABLED)

            # 统计信息
            stats_text = tk.Text(row_frame, font=font, width=80, height=row_height,
                                 wrap=tk.WORD, bg=stats_bg, fg=stats_fg, borderwidth=1, relief='solid')
            stats_text.pack(side=tk.LEFT, padx=5)
            stats_text.insert(tk.END, stats)
            stats_text.config(state=tk.DISABLED)

        # 更新Canvas滚动区域
        def update_scrollregion(event=None):
            canvas.configure(scrollregion=canvas.bbox("all"))
            canvas.itemconfig(canvas_window, width=canvas.winfo_width())

        inner_frame.bind("<Configure>", update_scrollregion)
        canvas.bind("<Configure>", lambda e: canvas.itemconfig(canvas_window, width=e.width))

        # 鼠标滚轮事件处理
        def on_mousewheel(event):
            canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

        def bind_mousewheel(event):
            canvas.bind_all("<MouseWheel>", on_mousewheel)

        def unbind_mousewheel(event):
            canvas.unbind_all("<MouseWheel>")

        canvas.bind('<Enter>', bind_mousewheel)
        canvas.bind('<Leave>', unbind_mousewheel)

        # 关闭按钮
        close_button = tk.Button(main_frame, text="关闭", command=stats_window.destroy,
                                 font=font, width=10, bg="#2196F3", fg="white")
        close_button.pack(pady=(10, 0))

    # 对比按钮
    compare_button = tk.Button(top_frame, text="对比EPLAN属性", command=start_compare, font=big_font, width=14, bg="#2196F3", fg="white")
    compare_button.pack(side=tk.LEFT, padx=5)

    # 下方Text区域
    text_frame = tk.Frame(root, padx=10, pady=10)
    text_frame.pack(fill=tk.BOTH, expand=True)

    # 使用PanedWindow创建可调整大小的左右区域
    paned_window = tk.PanedWindow(text_frame, orient=tk.HORIZONTAL)
    paned_window.pack(fill=tk.BOTH, expand=True)

    # 左边：识别属性Canvas
    left_frame = tk.Frame(paned_window)
    paned_window.add(left_frame, width=3 * root.winfo_width() // 4)

    # 识别Canvas
    tk.Label(left_frame, text="图纸取值(左) vs EPLAN属性(右)", font=("ABBvoice CNSG", 10, 'bold'), anchor='center').pack(fill=tk.X, pady=(0, 5))

    # Canvas框架
    canvas_frame = tk.Frame(left_frame)
    canvas_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 10))

    # 创建Canvas
    recognition_canvas = tk.Canvas(canvas_frame, bg='white')
    recognition_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

    # 滚动条
    scrollbar = ttk.Scrollbar(canvas_frame, orient=tk.VERTICAL, command=recognition_canvas.yview)
    scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
    recognition_canvas.configure(yscrollcommand=scrollbar.set)

    # 全局变量，用于存储对比值
    compare_values_cache = {}

    # 存储所有页面的数据
    all_pages_data = []

    # 内部框架，用于放置内容
    canvas_inner_frame = tk.Frame(recognition_canvas, bg='white')
    recognition_canvas.create_window((0, 0), window=canvas_inner_frame, anchor=tk.NW)

    # 显示当前页的对比结果
    def show_current_page_comparison():
        """
        显示当前页的对比结果
        """
        # 清空对比结果
        compare_text.config(state=tk.NORMAL)
        compare_text.delete(1.0, tk.END)
        compare_text.insert(tk.END, "当前页对比结果:\n")
        compare_text.insert(tk.END, "-" * 60 + "\n")

        # 检查是否有识别值和对比值
        if hasattr(canvas_inner_frame, 'recognized_value_texts') and hasattr(canvas_inner_frame, 'compare_value_texts'):
            differences_found = False

            # 按照属性的顺序进行对比，而不是随机顺序
            # 遍历recognized_value_texts中的属性，保持原有顺序
            for attr_name in canvas_inner_frame.recognized_value_texts:
                if attr_name in canvas_inner_frame.compare_value_texts:
                    # 获取识别值和对比值
                    recognized_text = canvas_inner_frame.recognized_value_texts[attr_name]
                    compare_text_widget = canvas_inner_frame.compare_value_texts[attr_name]

                    # 获取文本内容
                    try:
                        # 检查Text组件是否有效
                        recognized_text.winfo_exists()
                        compare_text_widget.winfo_exists()

                        recognized_value = recognized_text.get(1.0, tk.END).strip()
                        compare_value = compare_text_widget.get(1.0, tk.END).strip()

                        # 检查值是否只包含空格或为空
                        original_recognized = recognized_text.get(1.0, tk.END).rstrip('\n')
                        original_compare = compare_text_widget.get(1.0, tk.END).rstrip('\n')

                        # 如果值只包含空格或为空，保持原值或转换为空格
                        if original_recognized.strip() == "":
                            if original_recognized.isspace():
                                recognized_value = original_recognized
                            else:
                                # 为空字符串时转换为空格
                                recognized_value = " "
                        if original_compare.strip() == "":
                            if original_compare.isspace():
                                compare_value = original_compare
                            else:
                                # 为空字符串时转换为空格
                                compare_value = " "

                        # 对比值
                        if recognized_value != compare_value:
                            differences_found = True
                            compare_text.insert(tk.END, f"\n属性: {attr_name}\n")
                            compare_text.insert(tk.END, f"图纸识别值: {recognized_value}\n")
                            compare_text.insert(tk.END, f"EPLAN属性值: {compare_value}\n")
                            compare_text.insert(tk.END, "-" * 60 + "\n")
                            # 只对compare_value_text标红，不再对recognized_value_text标红
                            try:
                                compare_text_widget.config(state=tk.NORMAL)
                                compare_text_widget.delete(1.0, tk.END)
                                compare_text_widget.tag_configure("red", background="red", foreground="white")
                                # 保持原值，不使用strip后的值
                                if original_compare and not original_compare.isspace():
                                    compare_text_widget.insert(tk.END, original_compare, "red")
                                else:
                                    # 如果为空或只包含空格，插入空格占位
                                    compare_text_widget.insert(tk.END, "         ", "red")
                                compare_text_widget.config(state=tk.DISABLED)
                            except Exception as e:
                                # 不再打印错误信息，避免干扰用户
                                continue
                    except Exception as e:
                        # 不再打印错误信息，避免干扰用户
                        continue

            if not differences_found:
                compare_text.insert(tk.END, "无差异\n")
        else:
            compare_text.insert(tk.END, "无对比数据\n")

        compare_text.config(state=tk.DISABLED)

    # 导出Excel文件
    def is_consecutive(station1, station2):
        """
        检查两个站号是否连续
        支持格式：A01, A02...A99, B01...
        """
        match1 = re.match(r'^([A-Za-z])(\d{2})$', station1)
        match2 = re.match(r'^([A-Za-z])(\d{2})$', station2)

        if match1 and match2:
            letter1, num1 = match1.groups()
            letter2, num2 = match2.groups()

            if letter1 == letter2:
                return int(num2) == int(num1) + 1

        return False

    def generate_statistics(station_names, values_by_station):
        """
        生成统计信息，将具有相同值的站点分组
        返回格式：A01: XX\nA02-A08: YY
        返回：(统计字符串, 值的种类数量)
        """
        value_to_stations = {}
        for station in station_names:
            value = values_by_station[station]
            if value not in value_to_stations:
                value_to_stations[value] = []
            value_to_stations[value].append(station)

        result = []
        for value, stations in value_to_stations.items():
            if not stations:
                continue

            sorted_stations = sorted(stations)

            grouped_stations = []
            i = 0
            while i < len(sorted_stations):
                start = sorted_stations[i]
                end = start
                while i + 1 < len(sorted_stations):
                    next_station = sorted_stations[i + 1]
                    if is_consecutive(end, next_station):
                        end = next_station
                        i += 1
                    else:
                        break
                if start == end:
                    grouped_stations.append(start)
                else:
                    grouped_stations.append(f"{start}-{end}")
                i += 1

            stations_str = ", ".join(grouped_stations)
            result.append(f"{stations_str}: {value}")

        return "\n".join(result), len(value_to_stations)

    def create_summary_sheet(workbook, all_pages_data, thin_border):
        """
        创建Summary表单，汇总所有站点的数据
        第一列：属性
        第二列：统计（按分类汇总各站的识别值）
        第三列及以后：各站的识别值，列标题为站名
        """
        summary_sheet = workbook.create_sheet(title='Summary', index=0)

        all_attributes = []
        attr_set = set()
        for page_data in all_pages_data:
            for attr in page_data['attributes']:
                if attr['name'] not in attr_set:
                    attr_set.add(attr['name'])
                    all_attributes.append(attr['name'])

        station_names = [page_data['title'] for page_data in all_pages_data]

        summary_sheet['A1'] = "属性"
        summary_sheet['B1'] = "各站图纸取值对比"
        for col_idx, station_name in enumerate(station_names, 3):
            col_letter = chr(ord('A') + col_idx - 1)
            summary_sheet[f'{col_letter}1'] = station_name

        from openpyxl.styles import Alignment
        normal_font = styles.Font(name='ABBvoice CNSG')
        header_font = styles.Font(name='ABBvoice CNSG', bold=True)
        header_fill = styles.PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
        wrap_text_alignment = Alignment(wrap_text=True, vertical='top')

        for col in range(1, len(station_names) + 3):
            col_letter = chr(ord('A') + col - 1)
            summary_sheet[f'{col_letter}1'].font = header_font
            summary_sheet[f'{col_letter}1'].fill = header_fill
            summary_sheet[f'{col_letter}1'].border = thin_border
            summary_sheet[f'{col_letter}1'].alignment = wrap_text_alignment

        row = 2
        for attr_name in all_attributes:
            values_by_station = {}
            for page_data in all_pages_data:
                station_name = page_data['title']
                value = ""
                for attr in page_data['attributes']:
                    if attr['name'] == attr_name:
                        value = attr['recognized_value']
                        break
                values_by_station[station_name] = value if value.strip() else " "

            summary_sheet[f'A{row}'] = attr_name
            summary_sheet[f'A{row}'].border = thin_border
            summary_sheet[f'A{row}'].alignment = wrap_text_alignment
            summary_sheet[f'A{row}'].font = normal_font

            for col_idx, station_name in enumerate(station_names, 3):
                col_letter = chr(ord('A') + col_idx - 1)
                summary_sheet[f'{col_letter}{row}'] = values_by_station[station_name]
                summary_sheet[f'{col_letter}{row}'].border = thin_border
                summary_sheet[f'{col_letter}{row}'].alignment = wrap_text_alignment
                summary_sheet[f'{col_letter}{row}'].font = normal_font

            stats, value_count = generate_statistics(station_names, values_by_station)
            summary_sheet[f'B{row}'] = stats
            summary_sheet[f'B{row}'].border = thin_border
            summary_sheet[f'B{row}'].alignment = wrap_text_alignment
            summary_sheet[f'B{row}'].font = normal_font

            # 如果有多种值，标红
            if value_count > 1:
                red_fill = styles.PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                summary_sheet[f'A{row}'].fill = red_fill
                summary_sheet[f'B{row}'].fill = red_fill

            row += 1

        summary_sheet.column_dimensions['A'].width = 50
        summary_sheet.column_dimensions['B'].width = 80
        for col_idx in range(len(station_names)):
            col_letter = chr(ord('C') + col_idx)
            summary_sheet.column_dimensions[col_letter].width = 80

    def export_to_excel():
        """
        导出Excel文件，以站号为Sheet名，将canvas中的属性，识别值，对比值存到excel表单，有差异的属性及识别值，对比值标红
        """
        if not all_pages_data:
            messagebox.showerror("错误", "没有可导出的数据")
            return

        # 选择保存路径
        # 获取桌面路径
        desktop_path = os.path.join(os.path.expanduser("~"), "Desktop")

        # 获取项目名（从PDF路径中提取）
        project_name = "项目"
        pdf_path = pdf_path_var.get()
        if pdf_path:
            # 从PDF文件名中提取项目名（去除扩展名）
            project_name = os.path.splitext(os.path.basename(pdf_path))[0]

        # 构建默认文件名
        default_filename = f"{project_name}-DataSheet差异表"

        save_path = filedialog.asksaveasfilename(
            title="保存Excel文件",
            initialdir=desktop_path,
            initialfile=default_filename,
            defaultextension=".xlsx",
            filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*")]
        )

        if not save_path:
            return

        try:
            # 创建Excel文件
            workbook = Workbook()

            # 定义边框样式
            from openpyxl.styles import Border, Side, Alignment
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

            # 定义自动换行样式
            wrap_text_alignment = Alignment(wrap_text=True, vertical='top')

            # 定义字体样式为ABBvoice CNSG
            normal_font = styles.Font(name='ABBvoice CNSG')
            header_font = styles.Font(name='ABBvoice CNSG', bold=True)

            # 遍历所有页面
            for page_data in all_pages_data:
                # 创建Sheet，以站号为Sheet名
                sheet_name = page_data['title']
                # 限制Sheet名长度为31个字符
                if len(sheet_name) > 31:
                    sheet_name = sheet_name[:31]

                # 创建Sheet
                sheet = workbook.create_sheet(title=sheet_name)

                # 设置表头
                sheet['A1'] = "属性"
                sheet['B1'] = "图纸识别值"
                sheet['C1'] = "EPLAN属性值"

                # 设置表头样式
                header_fill = styles.PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")

                for col in ['A', 'B', 'C']:
                    sheet[f'{col}1'].font = header_font
                    sheet[f'{col}1'].fill = header_fill
                    sheet[f'{col}1'].border = thin_border
                    sheet[f'{col}1'].alignment = wrap_text_alignment

                # 填充数据
                row = 2
                # 从page_data中获取该页面的属性数据
                for attr in page_data['attributes']:
                    # 获取属性名、识别值和对比值
                    attr_name = attr['name']
                    recognized_value = attr['recognized_value']
                    compare_value = attr['compare_value']

                    # 处理空值
                    if recognized_value.strip() == "":
                        if recognized_value.isspace():
                            recognized_value = recognized_value
                        else:
                            recognized_value = " "
                    if compare_value.strip() == "":
                        if compare_value.isspace():
                            compare_value = compare_value
                        else:
                            compare_value = " "

                    # 写入数据
                    sheet[f'A{row}'] = attr_name
                    sheet[f'B{row}'] = recognized_value
                    sheet[f'C{row}'] = compare_value

                    # 设置边框、自动换行和字体
                    for col in ['A', 'B', 'C']:
                        sheet[f'{col}{row}'].border = thin_border
                        sheet[f'{col}{row}'].alignment = wrap_text_alignment
                        sheet[f'{col}{row}'].font = normal_font

                    # 检查是否有差异
                    recognized_value_stripped = str(recognized_value).strip()
                    compare_value_stripped = str(compare_value).strip()

                    if recognized_value_stripped != compare_value_stripped:
                        # 标红差异行
                        red_fill = styles.PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                        for col in ['A', 'B', 'C']:
                            sheet[f'{col}{row}'].fill = red_fill

                    row += 1

                # 调整列宽
                sheet.column_dimensions['A'].width = 50
                sheet.column_dimensions['B'].width = 80
                sheet.column_dimensions['C'].width = 80

            # 创建Summary表单
            create_summary_sheet(workbook, all_pages_data, thin_border)

            # 删除默认的Sheet
            if 'Sheet' in workbook.sheetnames:
                workbook.remove(workbook['Sheet'])

            # 保存文件
            workbook.save(save_path)

            messagebox.showinfo("成功", f"Excel文件已导出到：{save_path}")

        except Exception as e:
            messagebox.showerror("错误", f"导出Excel文件失败：{e}")
            print(f"导出Excel文件失败：{e}")

    # 导出按钮
    export_button = tk.Button(top_frame, text="导出报表", command=export_to_excel, font=big_font, width=8, bg="#FF9800", fg="white")
    export_button.pack(side=tk.LEFT, padx=5)

    # 设置按钮初始状态
    recognize_button.config(state=tk.DISABLED)
    compare_button.config(state=tk.DISABLED)
    export_button.config(state=tk.DISABLED)

    # 添加鼠标滚轮事件绑定
    def on_mousewheel(event):
        recognition_canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

    def bind_mousewheel(event):
        recognition_canvas.bind_all("<MouseWheel>", on_mousewheel)

    def unbind_mousewheel(event):
        recognition_canvas.unbind_all("<MouseWheel>")

    recognition_canvas.bind('<Enter>', bind_mousewheel)
    recognition_canvas.bind('<Leave>', unbind_mousewheel)

    # 分页控件
    pagination_frame = tk.Frame(left_frame)
    pagination_frame.pack(fill=tk.X, pady=(0, 5))

    # 创建居中框架
    center_frame = tk.Frame(pagination_frame)
    center_frame.pack(anchor=tk.CENTER)

    prev_button = tk.Button(center_frame, text="上一页", font=font, width=8)
    prev_button.pack(side=tk.LEFT, padx=5)
    page_label = tk.Label(center_frame, text="第 1 页，共 1 页", font=font)
    page_label.pack(side=tk.LEFT, padx=10)
    next_button = tk.Button(center_frame, text="下一页", font=font, width=8)
    next_button.pack(side=tk.LEFT, padx=5)

    # 右边：识别结果和对比结果Text
    right_frame = tk.Frame(paned_window)
    paned_window.add(right_frame, width=root.winfo_width() // 3)

    # 识别Text
    tk.Label(right_frame, text="识别结果", font=font, anchor='w').pack(fill=tk.X, pady=(0, 5))

    # 识别Text框架
    recognition_text_frame = tk.Frame(right_frame)
    recognition_text_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 10))

    recognition_text = tk.Text(recognition_text_frame, font=font, wrap=tk.WORD, state=tk.DISABLED)
    recognition_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

    scrollbar1 = tk.Scrollbar(recognition_text_frame, command=recognition_text.yview)
    scrollbar1.pack(side=tk.RIGHT, fill=tk.Y)
    recognition_text.config(yscrollcommand=scrollbar1.set)

    # 对比Text
    tk.Label(right_frame, text="对比结果", font=font, anchor='w').pack(fill=tk.X, pady=(0, 5))

    # 对比Text框架
    compare_text_frame = tk.Frame(right_frame)
    compare_text_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 10))

    compare_text = tk.Text(compare_text_frame, font=font, wrap=tk.WORD, state=tk.DISABLED)
    compare_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

    scrollbar2 = tk.Scrollbar(compare_text_frame, command=compare_text.yview)
    scrollbar2.pack(side=tk.RIGHT, fill=tk.Y)
    compare_text.config(yscrollcommand=scrollbar2.set)

    root.mainloop()


if __name__ == "__main__":
    create_gui()